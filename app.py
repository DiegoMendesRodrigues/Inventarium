# Sistema criado por Diego Mendes Rodrigues, utilizando desenvolvimento de software e inteligência artificial
from flask import Flask, render_template, request, redirect, url_for, flash, session, jsonify, abort, send_file
from flask_sqlalchemy import SQLAlchemy
from datetime import datetime, date, timedelta
from email.message import EmailMessage
from email.utils import formataddr
from io import BytesIO
from pathlib import Path
import hashlib
import re
import secrets
import smtplib
import string
from copy import copy
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from zipfile import ZipFile

app = Flask(__name__)
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///inventarium.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SECRET_KEY'] = 'inventarium_secret_key_2024'
app.json.sort_keys = False  # preserva a ordem dos campos do /resumo

NECESSIDADES_CHAVE_ACESSO = [
    'Alterar e visualizar alterações',
    'Visualizar alterações',
]

db = SQLAlchemy(app)


def gerar_codigo_interno(model, prefixo):
    codigos = [c for (c,) in model.query.with_entities(model.codigo_interno).filter(model.codigo_interno.startswith(prefixo + '-')).all()]
    sequencias = []
    for c in codigos:
        try:
            sequencias.append(int(c.split('-')[-1]))
        except ValueError:
            pass
    proximo = (max(sequencias) + 1) if sequencias else 1
    return f'{prefixo}-{proximo:03d}'


def ordenar_lista(model, colunas, padrao='codigo_interno'):
    ordem = request.args.get('ordem', padrao)
    direcao = request.args.get('direcao', 'asc')
    if ordem not in colunas:
        ordem = padrao
    if direcao not in ('asc', 'desc'):
        direcao = 'asc'
    coluna = colunas[ordem].desc() if direcao == 'desc' else colunas[ordem]
    return model.query.order_by(coluna).all(), ordem, direcao


def empresa_do_form():
    emp_id = request.form.get('empresa_id', '')
    return Empresa.query.get(int(emp_id)) if emp_id.isdigit() else None


@app.template_filter('mascara_cnpj')
def mascara_cnpj(valor):
    """Exibe o CNPJ oculto no formato 00.0*****-00 (somente listagens)."""
    if not valor or len(valor) < 8:
        return valor or ''
    return valor[:4] + '*****' + valor[-3:]


@app.template_filter('mascara_cpf_usuarios')
def mascara_cpf_usuarios(valor):
    """Exibe o CPF oculto no formato 000.******-00 (listagem de Usuários)."""
    if not valor or len(valor) < 8:
        return valor or ''
    return valor[:4] + '******' + valor[-3:]


@app.template_filter('mascara_cpf_admin')
def mascara_cpf_admin(valor):
    """Exibe o CPF oculto no formato 000.*****-00 (listagem de Administradores)."""
    if not valor or len(valor) < 8:
        return valor or ''
    return valor[:4] + '*****' + valor[-3:]


@app.template_filter('mascara_cpf_chaves')
def mascara_cpf_chaves(valor):
    """Exibe o CPF oculto no formato 000.******-00 (listagem de Chaves de Acesso)."""
    if not valor or len(valor) < 8:
        return valor or ''
    return valor[:4] + '******' + valor[-3:]


@app.context_processor
def empresas_para_formularios():
    return {'empresas': Empresa.query.order_by(Empresa.nome_fantasia.asc()).all()}


@app.context_processor
def perfil_do_usuario():
    return {'perfil_admin': session.get('administrador_tipo') == 'Administrador'}


def area_local_do_form():
    local_id = request.form.get('local_de_uso_id', '')
    return AreaLocal.query.get(int(local_id)) if local_id.isdigit() else None


def processar_equipamentos_software(valores):
    """Resolve os valores do multi-select 'Equipamento que está Instalado' de Softwares:
    'w-<id>'/'n-<id>'/'a-<id>' viram host + id de vinculação e 'nao-instalado' vira a marca de não instalado.
    Retorna (codigo_equipamento, ids_de_estacoes, ids_de_notebooks, ids_de_all_in_one)."""
    hosts, w_ids, n_ids, a_ids = [], [], [], []
    nao_instalado = False
    for valor in valores:
        if valor == 'nao-instalado':
            nao_instalado = True
            continue
        if '-' not in valor:
            continue
        tipo, _, id_str = valor.partition('-')
        if not id_str.isdigit():
            continue
        if tipo == 'w':
            w = Workstation.query.get(int(id_str))
            if w and w.id not in w_ids:
                w_ids.append(w.id)
                hosts.append(w.host)
        elif tipo == 'n':
            n = Notebook.query.get(int(id_str))
            if n and n.id not in n_ids:
                n_ids.append(n.id)
                hosts.append(n.host)
        elif tipo == 'a':
            a = EstacaoAllInOne.query.get(int(id_str))
            if a and a.id not in a_ids:
                a_ids.append(a.id)
                hosts.append(a.host)
    if nao_instalado and hosts:
        nao_instalado = False  # um equipamento escolhido vence a marca de não instalado
    codigo = 'Não está instalado' if nao_instalado else ', '.join(hosts)
    return codigo, w_ids, n_ids, a_ids


def sincronizar_vinculos_software(sw, w_ids, n_ids, a_ids):
    """Regrava as linhas N:N do software (tabela software_equipamentos) com os equipamentos escolhidos.
    Chamada após o flush do software (criação) ou antes do commit (edição)."""
    SoftwareEquipamento.query.filter_by(software_id=sw.id).delete()
    for wid in w_ids:
        db.session.add(SoftwareEquipamento(software_id=sw.id, workstation_id=wid))
    for nid in n_ids:
        db.session.add(SoftwareEquipamento(software_id=sw.id, notebook_id=nid))
    for aid in a_ids:
        db.session.add(SoftwareEquipamento(software_id=sw.id, estacao_all_in_one_id=aid))


def equipamentos_selecionados(sw):
    """Códigos 'w-<id>'/'n-<id>'/'a-<id>' (separados por vírgula) já vinculados ao software, para pré-marcá-los
    no multi-select do form; a opção 'nao-instalado' é marcada quando não há nenhum vínculo."""
    codigos = []
    for lin in SoftwareEquipamento.query.filter_by(software_id=sw.id).all():
        if lin.workstation_id:
            codigos.append(f'w-{lin.workstation_id}')
        elif lin.notebook_id:
            codigos.append(f'n-{lin.notebook_id}')
        elif lin.estacao_all_in_one_id:
            codigos.append(f'a-{lin.estacao_all_in_one_id}')
    if not codigos and sw.workstation_id:
        codigos.append(f'w-{sw.workstation_id}')
    elif not codigos and sw.notebook_id:
        codigos.append(f'n-{sw.notebook_id}')
    if not codigos and (sw.codigo_equipamento or '') == 'Não está instalado':
        codigos.append('nao-instalado')
    return ','.join(codigos)


def mapa_equipamentos():
    """Mapa 'w-<id>'/'n-<id>'/'s-<id>'/'a-<id>' -> host, usado nas listagens (Estação de Trabalho,
    All-in-One, Notebook ou Servidor)."""
    mapa = {f'w-{w.id}': w.host for w in Workstation.query.all()}
    mapa.update({f'n-{n.id}': n.host for n in Notebook.query.all()})
    mapa.update({f's-{s.id}': s.host for s in Servidor.query.all()})
    mapa.update({f'a-{a.id}': a.host for a in EstacaoAllInOne.query.all()})
    return mapa


def itens_vinculados_por_host(host, classe, chave):
    """Lista (crescente) de registros de `classe` cujo vínculo de equipamento casa com o `host` da Estação/
    Notebook/All-in-One, ordenada pela `chave` (ex.: Periferico.descricao, Software.nome). Só exibição no
    form; não grava nada. Para Periféricos aceita o valor cru do select ('w-<id>'/'n-<id>'/'a-<id>') ou host
    legado; para Softwares consulta também a tabela N:N software_equipamentos e casa hosts separados por vírgula."""
    if not host:
        return []
    mapa = mapa_equipamentos()
    ids_vinculados = set()
    if classe is Software:
        w = Workstation.query.filter_by(host=host).first()
        n = Notebook.query.filter_by(host=host).first()
        a = EstacaoAllInOne.query.filter_by(host=host).first()
        for lin in SoftwareEquipamento.query.all():
            if ((w and lin.workstation_id == w.id) or (n and lin.notebook_id == n.id)
                    or (a and lin.estacao_all_in_one_id == a.id)):
                ids_vinculados.add(lin.software_id)
    itens = []
    for it in classe.query.all():
        eq = it.codigo_equipamento or ''
        alvos = [mapa.get(v.strip(), v.strip()) for v in eq.split(',')]
        if it.id in ids_vinculados or host in alvos:
            itens.append(it)
    return sorted(itens, key=lambda it: (getattr(it, chave) or '').lower())


def vincular_monitores(valores, codigo_equipamento, al, usr, antigos=None):
    """Ao salvar a Estação/Notebook/Servidor, grava nos monitores escolhidos o vínculo com o
    equipamento, o Local de Uso e o Usuário do formulário. Valores 'Selecione' e 'Não utilizado' são ignorados.
    Se antigos for passado (pares novo/antigo, mesma ordem de valores) e um campo que tinha um monitor
    agora está vazio/'Selecione'/'Não utilizado', esse monitor é desvinculado: código de equipamento,
    Usuário e Local de Uso ficam limpos (somente se ainda estiver vinculado a este equipamento)."""
    for i, valor in enumerate(valores):
        mon_antigo = antigos[i] if antigos is not None and i < len(antigos) else None
        if not valor or valor in ('Selecione', 'Não utilizado'):
            if mon_antigo and mon_antigo not in ('Selecione', 'Não utilizado') and mon_antigo != valor:
                mon = Monitor.query.filter_by(codigo_interno=mon_antigo).first()
                if mon and mon.codigo_equipamento == codigo_equipamento:
                    mon.codigo_equipamento = ''
                    mon.local_de_uso = ''
                    mon.area_local_id = None
                    mon.usuario = ''
                    mon.usuario_id = None
            continue
        mon = Monitor.query.filter_by(codigo_interno=valor).first()
        if not mon:
            continue
        mon.codigo_equipamento = codigo_equipamento
        mon.local_de_uso = f'{al.area} - {al.local}' if al else ''
        mon.area_local_id = al.id if al else None
        mon.usuario = usr.nome_completo if usr else ''
        mon.usuario_id = usr.id if usr else None


@app.context_processor
def areas_locais_para_formularios():
    areas = AreaLocal.query.order_by(AreaLocal.area.asc(), AreaLocal.local.asc()).all()
    return {'areas_locais': [{'id': a.id, 'area': a.area, 'local': a.local, 'empresa_id': a.empresa_id} for a in areas]}


@app.context_processor
def usuarios_para_formularios():
    usrs = Usuario.query.order_by(Usuario.nome_completo.asc()).all()
    return {'usuarios_form': [{'id': u.id, 'nome_completo': u.nome_completo, 'empresa_id': u.empresa_id} for u in usrs]}


@app.context_processor
def monitores_para_formularios():
    mons = Monitor.query.order_by(Monitor.codigo_interno.asc()).all()
    return {'monitores_form': [{'codigo_interno': m.codigo_interno, 'marca': m.marca, 'tipo': m.tipo,
                                'tamanho': m.tamanho, 'empresa_id': m.empresa_id,
                                'equipamento': m.codigo_equipamento or ''} for m in mons]}


@app.context_processor
def equipamentos_para_formularios():
    equips = [{'id': w.id, 'host': w.host, 'empresa_id': w.empresa_id, 'tipo': 'w'} for w in Workstation.query.all()]
    equips += [{'id': n.id, 'host': n.host, 'empresa_id': n.empresa_id, 'tipo': 'n'} for n in Notebook.query.all()]
    equips += [{'id': a.id, 'host': a.host, 'empresa_id': a.empresa_id, 'tipo': 'a'}
               for a in EstacaoAllInOne.query.all()]
    equips.sort(key=lambda e: e['host'] or '')
    com_servidores = equips + [{'id': s.id, 'host': s.host, 'empresa_id': s.empresa_id, 'tipo': 's'}
                               for s in Servidor.query.all()]
    com_servidores.sort(key=lambda e: e['host'] or '')
    return {'equipamentos_form': equips, 'equipamentos_monitor_form': com_servidores}


UNIDADES_CAPACIDADE = {'KB': 1, 'MB': 2, 'GB': 3, 'TB': 4}


def chave_capacidade(nome):
    """Ordena capacidades pelo valor numérico e depois pela unidade (GB antes de TB)."""
    texto = (nome or '').strip()
    encontrado = re.search(r'(\d+(?:[.,]\d+)?)\s*(KB|MB|GB|TB)', texto, re.IGNORECASE)
    if encontrado:
        return (float(encontrado.group(1).replace(',', '.')), UNIDADES_CAPACIDADE[encontrado.group(2).upper()], texto.lower())
    numeros = re.findall(r'\d+', texto)
    return (float(numeros[0]) if numeros else 0.0, 0, texto.lower())


def md5(texto):
    """Hash MD5 da senha (definido pelo usuário; sem uso de biblioteca externa)."""
    return hashlib.md5(texto.encode('utf-8')).hexdigest()


def senha_aleatoria():
    """Gera senha com 12 caracteres e ao menos um de cada categoria exigida."""
    categorias = (
        string.ascii_uppercase,
        string.ascii_lowercase,
        string.digits,
        string.punctuation,
    )
    alfabeto = ''.join(categorias)
    caracteres = [secrets.choice(categoria) for categoria in categorias]
    caracteres.extend(secrets.choice(alfabeto) for _ in range(12 - len(caracteres)))
    secrets.SystemRandom().shuffle(caracteres)
    return ''.join(caracteres)


def configuracao_email_atual(configuracao_salva=None):
    configuracao = configuracao_salva or ConfiguracaoSistema.query.order_by(ConfiguracaoSistema.id.asc()).first()
    if configuracao is None:
        return None
    return {
        'nome_fantasia_empresa': configuracao.nome_fantasia_empresa,
        'url_acesso_sistema': configuracao.url_acesso_sistema,
        'servidor_smtp': configuracao.servidor_smtp,
        'porta_smtp': configuracao.porta_smtp,
        'usuario_autenticacao': configuracao.usuario_autenticacao,
        'senha_autenticacao': configuracao.senha_autenticacao,
        'nome_remetente': configuracao.nome_remetente,
        'assunto_esqueci_senha': configuracao.assunto_esqueci_senha,
        'mensagem_esqueci_senha': configuracao.mensagem_esqueci_senha,
        'assunto_redefinicao_senha': configuracao.assunto_redefinicao_senha,
        'mensagem_redefinicao_senha': configuracao.mensagem_redefinicao_senha,
        'assunto_software_assinatura': configuracao.assunto_software_assinatura,
        'mensagem_software_assinatura': configuracao.mensagem_software_assinatura,
    }


def gerar_chave_acesso():
    """Chave de acesso aleatória de 128 caracteres (maiúsculas, minúsculas e números)."""
    alfabeto = string.ascii_letters + string.digits
    return ''.join(secrets.choice(alfabeto) for _ in range(128))


def converter_senhas_para_md5():
    """Converte para MD5 as senhas dos administradores que ainda estiverem em texto puro."""
    alteradas = 0
    for adm in Administrador.query.all():
        senha = (adm.senha or '').strip()
        if senha and not re.fullmatch(r'[0-9a-f]{32}', senha):
            adm.senha = md5(senha)
            alteradas += 1
    if alteradas:
        db.session.commit()
        print(f'Senhas convertidas para MD5: {alteradas}')
    return alteradas


ASSUNTO_SOFTWARE_ASSINATURA_PADRAO = 'Software com Licenças por Assinatura - {Nome_Remetente}'
MENSAGEM_SOFTWARE_ASSINATURA_PADRAO = (
    'Prezado(a) {Usuario_Administrador},\n\n'
    'Seguem os Software com Licenças por Assinatura que irão expirar em breve:\n'
    '{Lista_Softwares_por_Assinatura}\n\n'
    '{Nome_Fantasia_Empresa_Gestora}\n'
    'Mensagem automática, não responder.'
)


def migrar_configuracoes_email_software():
    """Adiciona os modelos de e-mail para licenças por assinatura preservando configurações existentes."""
    with db.engine.begin() as conexao:
        colunas = {linha[1] for linha in conexao.exec_driver_sql('PRAGMA table_info(configuracoes_sistema)')}
        if 'assunto_software_assinatura' not in colunas:
            conexao.exec_driver_sql(
                "ALTER TABLE configuracoes_sistema ADD COLUMN assunto_software_assinatura VARCHAR(100) NOT NULL DEFAULT ''"
            )
            conexao.exec_driver_sql(
                'UPDATE configuracoes_sistema SET assunto_software_assinatura = ?',
                (ASSUNTO_SOFTWARE_ASSINATURA_PADRAO,),
            )
        if 'mensagem_software_assinatura' not in colunas:
            conexao.exec_driver_sql(
                "ALTER TABLE configuracoes_sistema ADD COLUMN mensagem_software_assinatura TEXT NOT NULL DEFAULT ''"
            )
            conexao.exec_driver_sql(
                'UPDATE configuracoes_sistema SET mensagem_software_assinatura = ?',
                (MENSAGEM_SOFTWARE_ASSINATURA_PADRAO,),
            )


def migrar_vinculos_softwares():
    """Cria as linhas N:N (software_equipamentos) dos softwares antigos a partir das FKs
    workstation_id/notebook_id, para preencher a tabela nova sem perder dados existentes.
    Idempotente: pula softwares que já possuem vínculo; seguro a cada início (como o MD5)."""
    criadas = 0
    for sw in Software.query.all():
        if SoftwareEquipamento.query.filter_by(software_id=sw.id).first():
            continue
        if sw.workstation_id:
            db.session.add(SoftwareEquipamento(software_id=sw.id, workstation_id=sw.workstation_id))
            criadas += 1
        if sw.notebook_id:
            db.session.add(SoftwareEquipamento(software_id=sw.id, notebook_id=sw.notebook_id))
            criadas += 1
    if criadas:
        db.session.commit()
        print(f'Vínculos de softwares criados: {criadas}')
    return criadas


def aplicar_placeholders_email(modelo, configuracao, adm, senha, placeholders_adicionais=None):
    valores = {
        '': 'Inventarium',
        'Nome_Remetente': configuracao['nome_remetente'],
        'Usuario_Senha_Nome': adm.nome_completo,
        'URL_Acesso_Sistema': configuracao['url_acesso_sistema'],
        'Usuario_Senha_User': adm.usuario,
        'Usuario_Senha_Pass': senha,
        'Usuario_Nova_Senha': senha,
        'Nome_Fantasia_Empresa_Gestora': configuracao['nome_fantasia_empresa'],
    }
    if placeholders_adicionais:
        valores.update(placeholders_adicionais)
    return re.sub(r'\{([^{}]*)\}', lambda match: str(valores.get(match.group(1), match.group(0))), modelo)


def enviar_mensagem_email_configurada(mensagem, configuracao):
    host = configuracao['servidor_smtp']
    porta = configuracao['porta_smtp']
    if porta in (465, 995):
        conexao = smtplib.SMTP_SSL(host, porta)
    else:
        conexao = smtplib.SMTP(host, porta)
    with conexao as servidor:
        if porta not in (465, 995):
            servidor.ehlo()
            if servidor.has_extn('starttls'):
                servidor.starttls()
                servidor.ehlo()
        servidor.login(configuracao['usuario_autenticacao'], configuracao['senha_autenticacao'])
        servidor.send_message(mensagem)


def enviar_email_senha(adm, senha, tipo='esqueci'):
    """Envia os dados de acesso do administrador para o e-mail cadastrado."""
    if tipo not in ('esqueci', 'redefinicao'):
        raise ValueError('Tipo de e-mail de senha inválido.')
    configuracao_salva = ConfiguracaoSistema.query.order_by(ConfiguracaoSistema.id.asc()).first()
    if configuracao_salva is None:
        raise RuntimeError('As configurações de e-mail não estão cadastradas.')
    configuracao = configuracao_email_atual(configuracao_salva)
    if tipo == 'redefinicao':
        assunto_modelo = configuracao['assunto_redefinicao_senha']
        corpo_modelo = configuracao['mensagem_redefinicao_senha']
    else:
        assunto_modelo = configuracao['assunto_esqueci_senha']
        corpo_modelo = configuracao['mensagem_esqueci_senha']
    assunto_email = aplicar_placeholders_email(assunto_modelo, configuracao, adm, senha)
    mensagem_email = aplicar_placeholders_email(corpo_modelo, configuracao, adm, senha)
    mensagem = EmailMessage()
    mensagem['From'] = formataddr((configuracao['nome_remetente'], configuracao['usuario_autenticacao']))
    mensagem['To'] = formataddr((adm.nome_completo, adm.email))
    mensagem['Subject'] = assunto_email
    mensagem.set_content(mensagem_email)
    enviar_mensagem_email_configurada(mensagem, configuracao)


def softwares_que_irao_expirar():
    hoje = date.today()
    limite = hoje + timedelta(days=60)
    softwares_irao_expirar = ''
    softwares = db.session.query(Software, Empresa.nome_fantasia) \
        .outerjoin(Empresa, Software.empresa_id == Empresa.id) \
        .filter(Software.validade_assinatura.isnot(None),
                Software.validade_assinatura >= hoje,
                Software.validade_assinatura <= limite) \
        .order_by(Empresa.nome_fantasia.asc(), Software.nome.asc(), Software.marca.asc()).all()
    for software, nome_empresa in softwares:
        linha = (f'• {software.validade_assinatura.strftime("%d/%m/%Y")} - {nome_empresa or ""} - '
                 f'{software.nome} - {software.marca}')
        softwares_irao_expirar = f'{softwares_irao_expirar}\n{linha}' if softwares_irao_expirar else linha
    return softwares_irao_expirar


def enviar_email_softwares_assinatura(adm, lista_softwares=None):
    configuracao_salva = ConfiguracaoSistema.query.order_by(ConfiguracaoSistema.id.asc()).first()
    if configuracao_salva is None:
        raise RuntimeError('As configurações de e-mail não estão cadastradas.')
    configuracao = configuracao_email_atual(configuracao_salva)
    if lista_softwares is None:
        lista_softwares = softwares_que_irao_expirar()
    placeholders = {
        'Usuario_Administrador': adm.nome_completo,
        'Lista_Softwares_por_Assinatura': lista_softwares,
    }
    assunto_email = aplicar_placeholders_email(
        configuracao['assunto_software_assinatura'], configuracao, adm, '', placeholders
    )
    mensagem_email = aplicar_placeholders_email(
        configuracao['mensagem_software_assinatura'], configuracao, adm, '', placeholders
    )
    mensagem = EmailMessage()
    mensagem['From'] = formataddr((configuracao['nome_remetente'], configuracao['usuario_autenticacao']))
    mensagem['To'] = formataddr((adm.nome_completo, adm.email))
    mensagem['Subject'] = assunto_email
    mensagem.set_content(mensagem_email)
    enviar_mensagem_email_configurada(mensagem, configuracao)


def opcoes_inf(modelo, numerica=False):
    """Valores de uma tabela inf_* para os <select> dos formulários — somente leitura."""
    nomes = [linha.nome for linha in modelo.query.all()]
    if numerica:
        nomes.sort(key=chave_capacidade)
    else:
        nomes.sort(key=lambda n: (n or '').lower())
    return nomes


def opcoes_workstation():
    """Valores das tabelas inf_* usados no formulário de Estações de Trabalho."""
    return {
        'tipos_memoria_ram': opcoes_inf(InfMemoriaRamTipo),
        'memorias_ram': opcoes_inf(InfMemoriaRamQtde, numerica=True),
        'discos': opcoes_inf(InfDiscoTamanho, numerica=True),
        'processadores': opcoes_inf(InfProcessadorDesktop),
        'placas_rede': opcoes_inf(InfPlacaRede),
        'placas_mae': opcoes_inf(InfPlacaMaeDesktop),
        'placas_de_video': opcoes_inf(InfPlacaDeVideo),
        'tamanhos_monitor': opcoes_inf(InfTamanhoMonitor, numerica=True),
        'sistemas_operacionais': opcoes_inf(InfSoDesktop),
    }


def opcoes_estacao_all_in_one():
    """Valores das tabelas inf_* usados no formulário de Estações All-in-One:
    mesmo das Estações de Trabalho, com 'Integrada' antes da lista de placas de vídeo."""
    opcoes = opcoes_workstation()
    opcoes['placas_de_video'] = ['Integrada'] + opcoes['placas_de_video']
    return opcoes


def opcoes_notebook():
    """Valores das tabelas inf_* usados no formulário de Notebooks."""
    return {
        'tipos_memoria_ram': opcoes_inf(InfMemoriaRamTipo),
        'memorias_ram': opcoes_inf(InfMemoriaRamQtde, numerica=True),
        'discos': opcoes_inf(InfDiscoTamanho, numerica=True),
        'processadores': opcoes_inf(InfProcessadorNotebook),
        'placas_rede': opcoes_inf(InfPlacaRede),
        'placas_mae': opcoes_inf(InfPlacaMaeNotebook),
        'placas_de_video': opcoes_inf(InfPlacaDeVideo),
        'tamanhos_monitor': opcoes_inf(InfTamanhoMonitor, numerica=True),
        'sistemas_operacionais': opcoes_inf(InfSoDesktop),
    }


def opcoes_monitor():
    """Valores das tabelas inf_* usados no formulário de Monitores."""
    return {'tamanhos_monitor': opcoes_inf(InfTamanhoMonitor, numerica=True)}


def opcoes_servidor():
    """Valores das tabelas inf_* usados no formulário de Servidores."""
    return {
        'tipos_memoria_ram': opcoes_inf(InfMemoriaRamTipo),
        'memorias_ram': opcoes_inf(InfMemoriaRamQtde, numerica=True),
        'discos': opcoes_inf(InfDiscoTamanho, numerica=True),
        'processadores': opcoes_inf(InfProcessadorServidor),
        'placas_rede': opcoes_inf(InfPlacaRede),
        'placas_mae': opcoes_inf(InfPlacaMaeServidor),
        'tamanhos_monitor': opcoes_inf(InfTamanhoMonitor, numerica=True),
        'sistemas_operacionais': opcoes_inf(InfSoServidor),
    }


def opcoes_impressora():
    """Valores das tabelas inf_* usados no formulário de Impressoras."""
    return {'marcas': opcoes_inf(InfImpressora)}


def opcoes_periferico():
    """Valores das tabelas inf_* usados no formulário de Periféricos."""
    return {
        'marcas': opcoes_inf(InfPeriferico),
        'perifericos_computador': opcoes_inf(InfPerifericoComputador),
    }


def opcoes_nobreak():
    """Valores das tabelas inf_* usados no formulário de No-Breaks."""
    return {'marcas': opcoes_inf(InfNoBreak)}


def opcoes_rack():
    """Valores das tabelas inf_* usados no formulário de Racks (Hacks)."""
    return {'marcas': opcoes_inf(InfHack)}


def opcoes_televisao():
    """Valores das tabelas inf_* usados no formulário de Televisões."""
    return {'marcas': opcoes_inf(InfTelevisaoMarca)}


def opcoes_tablet():
    """Valores das tabelas inf_* usados no formulário de Tablets."""
    return {'marcas': opcoes_inf(InfTabletsMarca)}


def opcoes_licencas_software():
    """Tipos de licença (inf_sofwares_licencas) usados no campo Tipo de
    Licenciamento do formulário de Softwares."""
    return {'licencas': opcoes_inf(InfSofwaresLicencas)}


# ==================== HARDWARE ====================

class Workstation(db.Model):
    __tablename__ = 'workstations'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    marca = db.Column(db.String(50), nullable=False)
    host = db.Column(db.String(150), nullable=False)
    tipo_memoria_ram = db.Column(db.String(50), nullable=False)
    memoria_ram = db.Column(db.String(50), nullable=False)
    disco1 = db.Column(db.String(50), nullable=False)
    disco2 = db.Column(db.String(50), nullable=False)
    processador = db.Column(db.String(100), nullable=False)
    placa_de_video = db.Column(db.String(50), nullable=False)
    placa_de_rede = db.Column(db.String(50), nullable=False)
    placa_de_rede_ethernet_2 = db.Column(db.String(50), nullable=False)
    placa_de_rede_wifi = db.Column(db.String(50), nullable=False)
    placa_mae = db.Column(db.String(50), nullable=False)
    sistema_operacional = db.Column(db.String(50), nullable=False)
    teclado = db.Column(db.String(50), nullable=False)
    mouse = db.Column(db.String(50), nullable=False)
    monitor1 = db.Column(db.String(50), nullable=False)
    monitor2 = db.Column(db.String(50), nullable=False)
    local_de_uso = db.Column(db.String(50))
    usuario = db.Column(db.String(150))
    usuario_id = db.Column(db.Integer, db.ForeignKey('usuarios.id'))
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class EstacaoAllInOne(db.Model):
    __tablename__ = 'estacoes_all_in_one'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    marca = db.Column(db.String(50), nullable=False)
    host = db.Column(db.String(150), nullable=False)
    tipo_memoria_ram = db.Column(db.String(50), nullable=False)
    memoria_ram = db.Column(db.String(50), nullable=False)
    disco1 = db.Column(db.String(50), nullable=False)
    disco2 = db.Column(db.String(50), nullable=False)
    processador = db.Column(db.String(100), nullable=False)
    placa_de_video = db.Column(db.String(50), nullable=False)
    placa_de_rede = db.Column(db.String(50), nullable=False)
    placa_de_rede_ethernet_2 = db.Column(db.String(50), nullable=False)
    placa_de_rede_wifi = db.Column(db.String(50), nullable=False)
    placa_mae = db.Column(db.String(50), nullable=False)
    sistema_operacional = db.Column(db.String(50), nullable=False)
    teclado = db.Column(db.String(50), nullable=False)
    mouse = db.Column(db.String(50), nullable=False)
    monitor1 = db.Column(db.String(50), nullable=False)
    monitor2 = db.Column(db.String(50), nullable=False)
    local_de_uso = db.Column(db.String(50))
    usuario = db.Column(db.String(150))
    usuario_id = db.Column(db.Integer, db.ForeignKey('usuarios.id'))
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class Notebook(db.Model):
    __tablename__ = 'notebooks'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    marca = db.Column(db.String(50), nullable=False)
    host = db.Column(db.String(150), nullable=False)
    tipo_memoria_ram = db.Column(db.String(50), nullable=False)
    memoria_ram = db.Column(db.String(50), nullable=False)
    disco1 = db.Column(db.String(50), nullable=False)
    disco2 = db.Column(db.String(50), nullable=False)
    processador = db.Column(db.String(100), nullable=False)
    placa_de_video = db.Column(db.String(50), nullable=False)
    placa_de_rede = db.Column(db.String(50), nullable=False)
    placa_de_rede_ethernet_2 = db.Column(db.String(50), nullable=False)
    placa_de_rede_wifi = db.Column(db.String(50), nullable=False)
    placa_mae = db.Column(db.String(50), nullable=False)
    sistema_operacional = db.Column(db.String(50), nullable=False)
    teclado = db.Column(db.String(50), nullable=False)
    mouse = db.Column(db.String(50), nullable=False)
    monitor1 = db.Column(db.String(50), nullable=False)
    monitor2 = db.Column(db.String(50), nullable=False)
    local_de_uso = db.Column(db.String(50))
    usuario = db.Column(db.String(150))
    usuario_id = db.Column(db.Integer, db.ForeignKey('usuarios.id'))
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class Monitor(db.Model):
    __tablename__ = 'monitores'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    marca = db.Column(db.String(50), nullable=False)
    tamanho = db.Column(db.String(50), nullable=False)
    tipo = db.Column(db.String(50), nullable=False)
    codigo_equipamento = db.Column(db.String(50))
    local_de_uso = db.Column(db.String(50))
    usuario = db.Column(db.String(150))
    usuario_id = db.Column(db.Integer, db.ForeignKey('usuarios.id'))
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class Periferico(db.Model):
    __tablename__ = 'perifericos'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    descricao = db.Column(db.String(100), nullable=False)
    marca = db.Column(db.String(50), nullable=False)
    perifericos_multimidia = db.Column(db.Text, nullable=False)
    codigo_equipamento = db.Column(db.String(50))
    local_de_uso = db.Column(db.String(50))
    usuario = db.Column(db.String(150))
    usuario_id = db.Column(db.Integer, db.ForeignKey('usuarios.id'))
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class Impressora(db.Model):
    __tablename__ = 'impressoras'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    marca = db.Column(db.String(50), nullable=False)
    tipo = db.Column(db.String(50), nullable=False)
    codigo_equipamento = db.Column(db.String(50))
    local_de_uso = db.Column(db.String(50))
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class Tablet(db.Model):
    __tablename__ = 'tablets'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    marca = db.Column(db.String(50), nullable=False)
    nome = db.Column(db.String(50), nullable=False)
    cor = db.Column(db.String(50), nullable=False)
    armazenamento = db.Column(db.String(50), nullable=False)
    memoria_ram = db.Column(db.String(50), nullable=False)
    resolucao_tela = db.Column(db.String(50), nullable=False)
    conectividade = db.Column(db.Text, nullable=False)
    sistema_operacional = db.Column(db.String(50), nullable=False)
    local_de_uso = db.Column(db.String(150))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    usuario = db.Column(db.String(150), nullable=False)
    usuario_id = db.Column(db.Integer, db.ForeignKey('usuarios.id'))
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


# ==================== INFRAESTRUTURA ====================

class Servidor(db.Model):
    __tablename__ = 'servidores'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    marca = db.Column(db.String(50), nullable=False)
    host = db.Column(db.String(150), nullable=False)
    tipo_memoria_ram = db.Column(db.String(50), nullable=False)
    memoria_ram = db.Column(db.String(50), nullable=False)
    disco1 = db.Column(db.String(50), nullable=False)
    disco2 = db.Column(db.String(50), nullable=False)
    disco3 = db.Column(db.String(50), nullable=False)
    disco4 = db.Column(db.String(50), nullable=False)
    processador = db.Column(db.String(100), nullable=False)
    placa_de_video = db.Column(db.String(50), nullable=False)
    placa_de_rede1 = db.Column(db.String(50), nullable=False)
    placa_de_rede2 = db.Column(db.String(50), nullable=False)
    placa_de_rede3 = db.Column(db.String(50), nullable=False)
    placa_de_rede4 = db.Column(db.String(50), nullable=False)
    placa_de_rede_wifi = db.Column(db.String(50), nullable=False)
    placa_mae = db.Column(db.String(50), nullable=False)
    sistema_operacional = db.Column(db.String(50), nullable=False)
    teclado = db.Column(db.String(50), nullable=False)
    mouse = db.Column(db.String(50), nullable=False)
    monitor1 = db.Column(db.String(50), nullable=False)
    local_de_uso = db.Column(db.String(50))
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class Projetor(db.Model):
    __tablename__ = 'projetores'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    marca = db.Column(db.String(50), nullable=False)
    tecnologia_imagem = db.Column(db.String(100), nullable=False)
    tipo = db.Column(db.String(100), nullable=False)
    resolucao = db.Column(db.String(100), nullable=False)
    conectividade = db.Column(db.String(100), nullable=False)
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    local_de_uso = db.Column(db.String(150))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class Televisao(db.Model):
    __tablename__ = 'televisoes'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    marca = db.Column(db.String(50), nullable=False)
    tipo_tecnologia_imagem = db.Column(db.String(50), nullable=False)
    resolucao = db.Column(db.String(50), nullable=False)
    conectividade_video = db.Column(db.Text, nullable=False)
    conectividade_rede = db.Column(db.String(50), nullable=False)
    local_de_uso = db.Column(db.String(150))
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class Modem(db.Model):
    __tablename__ = 'modens'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    marca = db.Column(db.String(50), nullable=False)
    host = db.Column(db.String(150), nullable=False)
    operadora = db.Column(db.String(50), nullable=False)
    velocidade_rede = db.Column(db.String(50), nullable=False)
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    local_de_uso = db.Column(db.String(150))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class Roteador(db.Model):
    __tablename__ = 'roteadores'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    host = db.Column(db.String(150))
    marca = db.Column(db.String(50), nullable=False)
    tipo = db.Column(db.String(100), nullable=False)
    portas = db.Column(db.String(50), nullable=False)
    velocidade_rede = db.Column(db.String(50), nullable=False)
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    local_de_uso = db.Column(db.String(150))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class Switch(db.Model):
    __tablename__ = 'switches'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    host = db.Column(db.String(150))
    marca = db.Column(db.String(50), nullable=False)
    tipo = db.Column(db.String(100), nullable=False)
    portas = db.Column(db.String(50), nullable=False)
    velocidade_rede = db.Column(db.String(50), nullable=False)
    local_de_uso = db.Column(db.String(150))
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class Firewall(db.Model):
    __tablename__ = 'firewalls'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    host = db.Column(db.String(150))
    marca = db.Column(db.String(100), nullable=False)
    portas_externas = db.Column(db.String(50), nullable=False)
    portas_internas = db.Column(db.String(50), nullable=False)
    velocidade_rede = db.Column(db.String(50), nullable=False)
    local_de_uso = db.Column(db.String(150))
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class AccessPoint(db.Model):
    __tablename__ = 'access_points'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    host = db.Column(db.String(150))
    marca = db.Column(db.String(100), nullable=False)
    tipo = db.Column(db.String(100), nullable=False)
    portas = db.Column(db.String(50), nullable=False)
    velocidade_rede = db.Column(db.String(50), nullable=False)
    local_de_uso = db.Column(db.String(150))
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class NoBreak(db.Model):
    __tablename__ = 'nobreaks'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    marca = db.Column(db.String(100), nullable=False)
    capacidade_va = db.Column(db.String(50), nullable=False)
    potencia_watts = db.Column(db.String(50), nullable=False)
    local_de_uso = db.Column(db.String(150))
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class Rack(db.Model):
    __tablename__ = 'racks'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_interno = db.Column(db.String(50), unique=True, nullable=False)
    marca = db.Column(db.String(50), nullable=False)
    cor = db.Column(db.String(50), nullable=False)
    modelo = db.Column(db.String(50), nullable=False)
    largura = db.Column(db.String(10), nullable=False)
    profundidade = db.Column(db.String(10), nullable=False)
    altura = db.Column(db.String(10), nullable=False)
    us = db.Column(db.String(4), nullable=False)
    peso_produto = db.Column(db.String(4), nullable=False)
    capacidade_carga = db.Column(db.String(4))
    fechadura = db.Column(db.String(10), nullable=False)
    classificacao_ip = db.Column(db.String(10))
    pes_niveladores = db.Column(db.String(10), nullable=False)
    local_de_uso = db.Column(db.String(150))
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


# ==================== SOFTWARE ====================

class Software(db.Model):
    __tablename__ = 'softwares'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    codigo_equipamento = db.Column(db.String(50), nullable=False)
    categoria = db.Column(db.String(50), nullable=False)
    marca = db.Column(db.String(50), nullable=False)
    nome = db.Column(db.String(50), nullable=False)
    descricao = db.Column(db.Text, nullable=False)
    tipo_licenciamento = db.Column(db.String(50), nullable=False)
    validade_assinatura = db.Column(db.Date)
    chave = db.Column(db.String(300))
    local_de_uso = db.Column(db.String(150))
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    area_local_id = db.Column(db.Integer, db.ForeignKey('areas_locais.id'))
    workstation_id = db.Column(db.Integer, db.ForeignKey('workstations.id'))
    notebook_id = db.Column(db.Integer, db.ForeignKey('notebooks.id'))
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


# Vínculo N:N entre Softwares e Estações/Notebooks/All-in-One (uma linha por par escolhido no multi-select).
# Tabela interna, sem CRUD próprio: gravada/excluída apenas pelas rotas de Softwares e pelas
# rotas de exclusão de Estação/Notebook/All-in-One.
class SoftwareEquipamento(db.Model):
    __tablename__ = 'software_equipamentos'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    software_id = db.Column(db.Integer, db.ForeignKey('softwares.id'), nullable=False)
    workstation_id = db.Column(db.Integer, db.ForeignKey('workstations.id'))
    notebook_id = db.Column(db.Integer, db.ForeignKey('notebooks.id'))
    estacao_all_in_one_id = db.Column(db.Integer, db.ForeignKey('estacoes_all_in_one.id'))


# ==================== EMPRESA ====================

class Empresa(db.Model):
    __tablename__ = 'empresas'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    razao_social = db.Column(db.String(150), nullable=False)
    nome_fantasia = db.Column(db.String(150), nullable=False)
    cnpj = db.Column(db.String(18), nullable=False)
    cep = db.Column(db.String(9), nullable=False)
    logradouro = db.Column(db.String(150), nullable=False)
    bairro = db.Column(db.String(50), nullable=False)
    cidade = db.Column(db.String(50), nullable=False)
    estado = db.Column(db.String(50), nullable=False)
    telefones = db.Column(db.String(50), nullable=False)
    email = db.Column(db.String(50), nullable=False)
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


# ==================== USUARIO ====================

class Usuario(db.Model):
    __tablename__ = 'usuarios'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    cpf = db.Column(db.String(14), nullable=False)
    empresa_que_trabalha = db.Column(db.String(150), nullable=False)
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    nome_completo = db.Column(db.String(150), nullable=False)
    cep = db.Column(db.String(9), nullable=False)
    logradouro = db.Column(db.String(150), nullable=False)
    bairro = db.Column(db.String(50), nullable=False)
    cidade = db.Column(db.String(50), nullable=False)
    estado = db.Column(db.String(50), nullable=False)
    telefones = db.Column(db.String(50), nullable=False)
    email = db.Column(db.String(50), nullable=False)
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class Administrador(db.Model):
    __tablename__ = 'administradores'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    cpf = db.Column(db.String(15), nullable=False)
    nome_completo = db.Column(db.String(150), nullable=False)
    telefone = db.Column(db.String(20), nullable=False)
    email = db.Column(db.String(150), nullable=False)
    usuario = db.Column(db.String(50), nullable=False)
    tipo = db.Column(db.String(50), nullable=False)
    recebe_notificacoes = db.Column(db.String(3), nullable=False)
    senha = db.Column(db.String(50), nullable=False)
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


# ==================== AREA LOCAL ====================

class AreaLocal(db.Model):
    __tablename__ = 'areas_locais'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    empresa = db.Column(db.String(150), nullable=False)
    empresa_id = db.Column(db.Integer, db.ForeignKey('empresas.id'))
    area = db.Column(db.String(50), nullable=False)
    local = db.Column(db.String(50), nullable=False)
    outras_informacoes = db.Column(db.Text)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


# ==================== CHAVE DE ACESSO ====================

class ChaveAcesso(db.Model):
    __tablename__ = 'chaves_acesso'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    cpf = db.Column(db.String(14), nullable=False)
    nome_completo = db.Column(db.String(150), nullable=False)
    telefone = db.Column(db.String(16), nullable=False)
    email = db.Column(db.String(50), nullable=False)
    cep = db.Column(db.String(9), nullable=False)
    logradouro = db.Column(db.String(150), nullable=False)
    bairro = db.Column(db.String(50), nullable=False)
    cidade = db.Column(db.String(50), nullable=False)
    estado = db.Column(db.String(50), nullable=False)
    necessidade = db.Column(db.String(50), nullable=False, default='')
    necessidade_webservices = db.Column(db.Text, nullable=False)
    chave_acesso = db.Column(db.String(128), unique=True, nullable=False)
    data_cadastro = db.Column(db.DateTime, default=datetime.utcnow)


class ConfiguracaoSistema(db.Model):
    __tablename__ = 'configuracoes_sistema'
    id = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome_fantasia_empresa = db.Column(db.String(100), nullable=False)
    url_acesso_sistema = db.Column(db.String(200), nullable=False)
    senha_inicial_administrador = db.Column(db.String(50), nullable=False)
    servidor_smtp = db.Column(db.String(200), nullable=False)
    porta_smtp = db.Column(db.Integer, nullable=False)
    usuario_autenticacao = db.Column(db.String(200), nullable=False)
    senha_autenticacao = db.Column(db.String(200), nullable=False)
    nome_remetente = db.Column(db.String(100), nullable=False)
    assunto_esqueci_senha = db.Column(db.String(100), nullable=False)
    mensagem_esqueci_senha = db.Column(db.Text, nullable=False)
    assunto_redefinicao_senha = db.Column(db.String(100), nullable=False)
    mensagem_redefinicao_senha = db.Column(db.Text, nullable=False)
    assunto_software_assinatura = db.Column(db.String(100), nullable=False, default=ASSUNTO_SOFTWARE_ASSINATURA_PADRAO)
    mensagem_software_assinatura = db.Column(db.Text, nullable=False, default=MENSAGEM_SOFTWARE_ASSINATURA_PADRAO)


# ==================== TABELAS INF (dados de referência) ====================
# As tabelas abaixo alimentam os campos de referência da aplicação. Seus valores
# podem ser gerenciados por Administradores em Gestão > Configurações ou pela
# API autenticada /informacoes. O sistema não semeia nem restaura esses dados.


class InfProcessadorDesktop(db.Model):
    __tablename__ = 'inf_processadores_desktop'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfProcessadorNotebook(db.Model):
    __tablename__ = 'inf_processadores_notebook'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfProcessadorServidor(db.Model):
    __tablename__ = 'inf_processadores_servidor'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfMemoriaRamTipo(db.Model):
    __tablename__ = 'inf_memorias_ram_tipo'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfMemoriaRamQtde(db.Model):
    __tablename__ = 'inf_memorias_ram_qtde'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfDiscoTamanho(db.Model):
    __tablename__ = 'inf_disco_tamanho'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfPlacaMaeDesktop(db.Model):
    __tablename__ = 'inf_placa_mae_desktop'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfPlacaMaeNotebook(db.Model):
    __tablename__ = 'inf_placa_mae_notebook'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfPlacaMaeServidor(db.Model):
    __tablename__ = 'inf_placa_mae_servidor'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfTamanhoMonitor(db.Model):
    __tablename__ = 'inf_tamanho_monitor'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfPlacaRede(db.Model):
    __tablename__ = 'inf_placa_rede'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfPlacaDeVideo(db.Model):
    __tablename__ = 'inf_placas_de_video'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfImpressora(db.Model):
    __tablename__ = 'inf_impressoras'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfPeriferico(db.Model):
    __tablename__ = 'inf_perifericos'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfPerifericoComputador(db.Model):
    __tablename__ = 'inf_perifericos_computador'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfServidor(db.Model):
    __tablename__ = 'inf_servidores'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfNoBreak(db.Model):
    __tablename__ = 'inf_nobreaks'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfHack(db.Model):
    __tablename__ = 'inf_hacks'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfSoDesktop(db.Model):
    __tablename__ = 'inf_so_desktop'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfSoServidor(db.Model):
    __tablename__ = 'inf_so_servidores'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfTelevisaoMarca(db.Model):
    __tablename__ = 'inf_televisao_marca'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfTabletsMarca(db.Model):
    __tablename__ = 'inf_tablets_marcas'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


class InfSofwaresLicencas(db.Model):
    __tablename__ = 'inf_sofwares_licencas'
    codigo = db.Column(db.Integer, primary_key=True, autoincrement=True)
    nome = db.Column(db.String(100), nullable=False, index=True)


CONFIGURACOES_REFERENCIA = [
    {'rotulo': 'Discos - Tamanhos', 'tabela': 'inf_disco_tamanho', 'modelo': InfDiscoTamanho},
    {'rotulo': 'Marcas de Hacks', 'tabela': 'inf_hacks', 'modelo': InfHack},
    {'rotulo': 'Marcas de Impressoras', 'tabela': 'inf_impressoras', 'modelo': InfImpressora},
    {'rotulo': 'Memórias RAM - Quantidade', 'tabela': 'inf_memorias_ram_qtde', 'modelo': InfMemoriaRamQtde},
    {'rotulo': 'Memórias RAM - Tipos', 'tabela': 'inf_memorias_ram_tipo', 'modelo': InfMemoriaRamTipo},
    {'rotulo': 'Marcas de No-Breaks', 'tabela': 'inf_nobreaks', 'modelo': InfNoBreak},
    {'rotulo': 'Marcas de Periféricos', 'tabela': 'inf_perifericos', 'modelo': InfPeriferico},
    {'rotulo': 'Tipos de Periféricos', 'tabela': 'inf_perifericos_computador', 'modelo': InfPerifericoComputador},
    {'rotulo': 'Marcas de Placas-mãe de Desktops', 'tabela': 'inf_placa_mae_desktop', 'modelo': InfPlacaMaeDesktop},
    {'rotulo': 'Marcas de Placas-mãe de Notebooks', 'tabela': 'inf_placa_mae_notebook', 'modelo': InfPlacaMaeNotebook},
    {'rotulo': 'Marcas de Placas-mãe de Servidores', 'tabela': 'inf_placa_mae_servidor', 'modelo': InfPlacaMaeServidor},
    {'rotulo': 'Marcas de Placas de Redes', 'tabela': 'inf_placa_rede', 'modelo': InfPlacaRede},
    {'rotulo': 'Modelos de Placas de Vídeo', 'tabela': 'inf_placas_de_video', 'modelo': InfPlacaDeVideo},
    {'rotulo': 'Modelos de Processadores de Desktops', 'tabela': 'inf_processadores_desktop', 'modelo': InfProcessadorDesktop},
    {'rotulo': 'Modelos de Processadores de Notebooks', 'tabela': 'inf_processadores_notebook', 'modelo': InfProcessadorNotebook},
    {'rotulo': 'Modelos de Processadores de Servidores', 'tabela': 'inf_processadores_servidor', 'modelo': InfProcessadorServidor},
    {'rotulo': 'Marcas de Servidores', 'tabela': 'inf_servidores', 'modelo': InfServidor},
    {'rotulo': 'Sistemas Operacionais - Desktops', 'tabela': 'inf_so_desktop', 'modelo': InfSoDesktop},
    {'rotulo': 'Sistemas Operacionais - Servidores', 'tabela': 'inf_so_servidores', 'modelo': InfSoServidor},
    {'rotulo': 'Tipos de Licenças de Softwares', 'tabela': 'inf_sofwares_licencas', 'modelo': InfSofwaresLicencas},
    {'rotulo': 'Marcas de Tablets', 'tabela': 'inf_tablets_marcas', 'modelo': InfTabletsMarca},
    {'rotulo': 'Monitores - Tamanhos', 'tabela': 'inf_tamanho_monitor', 'modelo': InfTamanhoMonitor},
    {'rotulo': 'Marcas de Televisões', 'tabela': 'inf_televisao_marca', 'modelo': InfTelevisaoMarca},
]
TABELAS_CONFIGURACOES = {item['tabela']: item for item in CONFIGURACOES_REFERENCIA}

REFERENCIAS_CADASTROS = {
    'inf_disco_tamanho': [
        (Workstation, ('disco1', 'disco2')),
        (EstacaoAllInOne, ('disco1', 'disco2')),
        (Notebook, ('disco1', 'disco2')),
        (Servidor, ('disco1', 'disco2', 'disco3', 'disco4')),
    ],
    'inf_memorias_ram_tipo': [
        (Workstation, ('tipo_memoria_ram',)),
        (EstacaoAllInOne, ('tipo_memoria_ram',)),
        (Notebook, ('tipo_memoria_ram',)),
        (Servidor, ('tipo_memoria_ram',)),
    ],
    'inf_memorias_ram_qtde': [
        (Workstation, ('memoria_ram',)),
        (EstacaoAllInOne, ('memoria_ram',)),
        (Notebook, ('memoria_ram',)),
        (Servidor, ('memoria_ram',)),
    ],
    'inf_processadores_desktop': [
        (Workstation, ('processador',)),
        (EstacaoAllInOne, ('processador',)),
    ],
    'inf_processadores_notebook': [(Notebook, ('processador',))],
    'inf_processadores_servidor': [(Servidor, ('processador',))],
    'inf_placa_mae_desktop': [
        (Workstation, ('placa_mae',)),
        (EstacaoAllInOne, ('placa_mae',)),
    ],
    'inf_placa_mae_notebook': [(Notebook, ('placa_mae',))],
    'inf_placa_mae_servidor': [(Servidor, ('placa_mae',))],
    'inf_tamanho_monitor': [
        (Monitor, ('tamanho',)),
        (EstacaoAllInOne, ('monitor1',)),
        (Notebook, ('monitor1',)),
    ],
    'inf_placa_rede': [
        (Workstation, ('placa_de_rede', 'placa_de_rede_ethernet_2', 'placa_de_rede_wifi')),
        (EstacaoAllInOne, ('placa_de_rede', 'placa_de_rede_ethernet_2', 'placa_de_rede_wifi')),
        (Notebook, ('placa_de_rede', 'placa_de_rede_ethernet_2', 'placa_de_rede_wifi')),
        (Servidor, ('placa_de_rede1', 'placa_de_rede2', 'placa_de_rede3', 'placa_de_rede4', 'placa_de_rede_wifi')),
    ],
    'inf_impressoras': [(Impressora, ('marca',))],
    'inf_perifericos': [(Periferico, ('marca',))],
    'inf_perifericos_computador': [(Periferico, ('perifericos_multimidia',))],
    'inf_servidores': [(Servidor, ('marca',))],
    'inf_nobreaks': [(NoBreak, ('marca',))],
    'inf_hacks': [(Rack, ('marca',))],
    'inf_so_desktop': [
        (Workstation, ('sistema_operacional',)),
        (EstacaoAllInOne, ('sistema_operacional',)),
        (Notebook, ('sistema_operacional',)),
    ],
    'inf_so_servidores': [(Servidor, ('sistema_operacional',))],
    'inf_televisao_marca': [(Televisao, ('marca',))],
    'inf_placas_de_video': [
        (Workstation, ('placa_de_video',)),
        (EstacaoAllInOne, ('placa_de_video',)),
        (Notebook, ('placa_de_video',)),
    ],
    'inf_tablets_marcas': [(Tablet, ('marca',))],
    'inf_sofwares_licencas': [(Software, ('tipo_licenciamento',))],
}
REFERENCIAS_CADASTROS_MULTIVALOR = {
    'inf_perifericos_computador': [(Periferico, 'perifericos_multimidia')],
}


def valores_referencia_em_uso(tabela):
    valores = set()
    for modelo, campos in REFERENCIAS_CADASTROS.get(tabela, []):
        for campo in campos:
            coluna = getattr(modelo, campo)
            for (valor,) in db.session.query(coluna).filter(coluna.isnot(None)).distinct().all():
                if valor and valor.strip():
                    valores.add(valor.strip())
    for modelo, campo in REFERENCIAS_CADASTROS_MULTIVALOR.get(tabela, []):
        coluna = getattr(modelo, campo)
        for (valor,) in db.session.query(coluna).filter(coluna.isnot(None)).distinct().all():
            if valor:
                valores.update(parte.strip() for parte in valor.split(',') if parte.strip())
    return valores


@app.context_processor
def configuracoes_para_menu():
    return {'configuracoes_referencia': CONFIGURACOES_REFERENCIA}


# ==================== ROTAS ====================

# ---------- Autenticação ----------

@app.before_request
def exigir_login():
    """Sem sessão autenticada, todas as telas redirecionam para a Autenticação.
    Perfis diferentes de 'Administrador' apenas visualizam: bloqueia inclusão
    (telas novo/nova) e qualquer gravação (POST); as telas de editar abrem em
    modo de leitura, com campos e Salvar desabilitados pelo JS do base.html."""
    if request.endpoint in ('autenticar', 'ajuda', 'esqueci_senha', 'sair', 'static') or request.path in ('/resumo', '/ativos', '/informacoes'):
        return None
    if 'administrador_id' not in session:
        return redirect(url_for('autenticar'))
    if session.get('administrador_tipo') != 'Administrador':
        if request.method not in ('GET', 'HEAD', 'OPTIONS'):
            flash('Acesso restrito: seu perfil permite apenas visualizar as informações.', 'error')
            return redirect(url_for('dashboard'))
        if request.endpoint and request.endpoint.startswith(('novo_', 'nova_')):
            flash('Acesso restrito: seu perfil permite apenas visualizar as informações.', 'error')
            return redirect(url_for('dashboard'))
    return None


@app.route('/autenticar', methods=['GET', 'POST'])
def autenticar():
    if 'administrador_id' in session:
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        usuario = request.form.get('usuario', '').strip()
        senha = request.form.get('senha', '')
        adm = Administrador.query.filter_by(usuario=usuario).first()
        if adm and adm.senha == md5(senha):
            session['administrador_id'] = adm.id
            session['administrador_nome'] = adm.nome_completo
            session['administrador_tipo'] = adm.tipo
            return redirect(url_for('dashboard'))
        return render_template('autenticacao.html', erro='Usuário e/ou senha inválidos')
    return render_template('autenticacao.html')


@app.route('/ajuda')
def ajuda():
    return render_template('ajuda.html')


@app.route('/sair', methods=['POST'])
def sair():
    session.clear()
    return redirect(url_for('autenticar'))


@app.route('/alterar-senha', methods=['POST'])
def alterar_senha():
    ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
    adm = Administrador.query.get(session.get('administrador_id')) if 'administrador_id' in session else None
    if not adm:
        if ajax:
            return jsonify(ok=False, erro='Sessão expirada. Autentique-se novamente.')
        return redirect(url_for('autenticar'))
    atual = request.form.get('senha_atual', '')
    nova = request.form.get('nova_senha', '')
    confirme = request.form.get('confirme_senha', '')
    faltando = []
    if not atual:
        faltando.append('Senha Atual')
    if not nova:
        faltando.append('Nova senha')
    if not confirme:
        faltando.append('Confirme a Nova senha')
    if faltando:
        erro = 'Preencha os campos: ' + ', '.join(faltando) + '.'
    elif len(atual) < 8 or len(nova) < 8 or len(confirme) < 8:
        erro = 'As senhas precisam possuir no mínimo 8 caracteres.'
    elif nova != confirme:
        erro = 'O campo Nova senha e o campo Confirme a Nova senha não são iguais.'
    elif adm.senha != md5(atual):
        erro = 'A Senha Atual está inválida.'
    else:
        adm.senha = md5(nova)
        db.session.commit()
        session.pop('administrador_id', None)
        session.pop('administrador_nome', None)
        session.pop('administrador_tipo', None)
        if ajax:
            return jsonify(ok=True)
        flash('Senha Alterada com Sucesso!', 'success')
        return redirect(url_for('autenticar'))
    if ajax:
        return jsonify(ok=False, erro=erro)
    flash(erro, 'error')
    return redirect(request.referrer or url_for('dashboard'))


@app.route('/esqueci-senha', methods=['GET', 'POST'])
def esqueci_senha():
    if request.method == 'POST':
        email = request.form.get('email', '').strip()
        adm = Administrador.query.filter_by(email=email).first() if email else None
        smtp = ConfiguracaoSistema.query.order_by(ConfiguracaoSistema.id.asc()).first()
        if adm and smtp and smtp.usuario_autenticacao and smtp.senha_autenticacao:
            nova_senha = senha_aleatoria()
            adm.senha = md5(nova_senha)
            db.session.commit()
            try:
                enviar_email_senha(adm, nova_senha, tipo='esqueci')
            except Exception as erro:
                print(f'Erro ao enviar e-mail de recuperação: {erro}')
            try:
                enviar_email_softwares_assinatura(adm)
            except Exception as erro:
                print(f'Erro ao enviar e-mail de softwares por assinatura: {erro}')
        return render_template('esqueci_senha.html', notificacao=True)
    return render_template('esqueci_senha.html')


@app.route('/')
def dashboard():
    stats = {
        'workstations': Workstation.query.count(),
        'estacoes_all_in_one': EstacaoAllInOne.query.count(),
        'notebooks': Notebook.query.count(),
        'monitores': Monitor.query.count(),
        'perifericos': Periferico.query.count(),
        'impressoras': Impressora.query.count(),
        'tablets': Tablet.query.count(),
        'servidores': Servidor.query.count(),
        'modens': Modem.query.count(),
        'roteadores': Roteador.query.count(),
        'switches': Switch.query.count(),
        'firewalls': Firewall.query.count(),
        'access_points': AccessPoint.query.count(),
        'projetores': Projetor.query.count(),
        'televisoes': Televisao.query.count(),
        'nobreaks': NoBreak.query.count(),
        'racks': Rack.query.count(),
        'softwares': Software.query.count(),
    }
    modelos_ativos = [Workstation, EstacaoAllInOne, Notebook, Monitor, Periferico, Impressora, Tablet, Servidor,
                      Modem, Firewall, Roteador, Switch, AccessPoint, Projetor, NoBreak, Rack]
    contagens = {}
    for modelo in modelos_ativos:
        for empresa_id, qtde in db.session.query(modelo.empresa_id, db.func.count(modelo.id)).group_by(modelo.empresa_id).all():
            if empresa_id is not None:
                contagens[empresa_id] = contagens.get(empresa_id, 0) + qtde
    consulta_softwares = db.session.query(Software.empresa_id, db.func.count(Software.id))
    for empresa_id, qtde in consulta_softwares.group_by(Software.empresa_id).all():
        if empresa_id is not None:
            contagens[empresa_id] = contagens.get(empresa_id, 0) + qtde
    empresas_stats = [(e.nome_fantasia, contagens.get(e.id, 0))
                      for e in Empresa.query.order_by(Empresa.nome_fantasia.asc()).all()]
    total = sum(qtde for _, qtde in empresas_stats)
    tipo_licenciamento_normalizado = db.func.lower(Software.tipo_licenciamento)
    softwares_perpetuos = Software.query.filter(
        ~tipo_licenciamento_normalizado.contains('assinatura')
    ) \
        .order_by(Software.nome.asc()).all()
    softwares_assinatura = Software.query.filter(
        tipo_licenciamento_normalizado.contains('assinatura')
    ) \
        .order_by(Software.validade_assinatura.asc(), Software.nome.asc()).all()
    return render_template('dashboard.html', stats=stats, total=total, empresas_stats=empresas_stats,
                           softwares_perpetuos=softwares_perpetuos, softwares_assinatura=softwares_assinatura)


# ---------- Webservice REST ----------

def exigir_bearer():
    """Webservices /resumo, /ativos e /informacoes exigem o cabeçalho
    Authorization: Bearer <chave de acesso>, cujo valor deve existir no
    cadastro Chaves de Acesso. A Necessidade da chave limita os métodos
    permitidos; retorna 401 para chave ausente/inválida e 403 sem permissão."""
    partes = request.headers.get('Authorization', '').split()
    chave = None
    if len(partes) == 2 and partes[0].lower() == 'bearer':
        chave = ChaveAcesso.query.filter_by(chave_acesso=partes[1]).first()
    if chave is None:
        return jsonify({'erro': 'Autenticação Bearer Token inválida ou ausente'}), 401, \
            {'WWW-Authenticate': 'Bearer'}
    if request.method == 'GET':
        if chave.necessidade in NECESSIDADES_CHAVE_ACESSO:
            return None
    elif request.method in ('POST', 'PUT') \
            and chave.necessidade == 'Alterar e visualizar alterações':
        return None
    return jsonify({'erro': 'A chave de acesso não possui permissão para esta requisição.'}), 403


@app.route('/resumo')
def resumo():
    """Webservice com Bearer Token (Chaves de Acesso): ativos por empresa em JSON."""
    negado = exigir_bearer()
    if negado:
        return negado
    def contar(modelo, grupo, chave, filtro=None):
        consulta = db.session.query(modelo.empresa_id, db.func.count(modelo.id))
        if filtro is not None:
            consulta = consulta.filter(modelo.tipo_licenciamento == filtro)
        for empresa_id, qtde in consulta.group_by(modelo.empresa_id).all():
            if empresa_id is None:
                continue
            bloco = contagens.setdefault(empresa_id, {}).setdefault(grupo, {})
            bloco[chave] = bloco.get(chave, 0) + qtde

    contagens = {}
    contar(Workstation, 'Hardware', 'Estacoes-de-Trabalho')
    contar(EstacaoAllInOne, 'Hardware', 'Estacoes-All-in-One')
    contar(Notebook, 'Hardware', 'Notebooks')
    contar(Monitor, 'Hardware', 'Monitores')
    contar(Impressora, 'Hardware', 'Impressoras')
    contar(Periferico, 'Hardware', 'Perifericos')
    contar(Tablet, 'Hardware', 'Tablets')
    contar(Modem, 'Infraestrutura', 'Modens')
    contar(Firewall, 'Infraestrutura', 'Firewalls')
    contar(Roteador, 'Infraestrutura', 'Roteadores')
    contar(Switch, 'Infraestrutura', 'Switches')
    contar(AccessPoint, 'Infraestrutura', 'Access-Points')
    contar(Servidor, 'Infraestrutura', 'Servidores')
    contar(Projetor, 'Infraestrutura', 'Projetores')
    contar(Televisao, 'Infraestrutura', 'Televisoes')
    contar(NoBreak, 'Infraestrutura', 'No-Breaks')
    contar(Rack, 'Infraestrutura', 'Racks')
    contar(Software, 'Software', 'Softwares-Licenciamento-Perpetuo', filtro='Perpétuo')
    contar(Software, 'Software', 'Softwares-Por-Assinatura', filtro='Por assinatura')

    modelo_vazio = {'Hardware': {'Estacoes-de-Trabalho': 0, 'Estacoes-All-in-One': 0, 'Notebooks': 0,
                                 'Monitores': 0, 'Impressoras': 0, 'Perifericos': 0, 'Tablets': 0},
                    'Infraestrutura': {'Modens': 0, 'Firewalls': 0, 'Roteadores': 0, 'Switches': 0,
                                       'Access-Points': 0, 'Servidores': 0, 'Projetores': 0,
                                       'Televisoes': 0, 'No-Breaks': 0, 'Racks': 0},
                    'Software': {'Softwares-Licenciamento-Perpetuo': 0,
                                 'Softwares-Por-Assinatura': 0}}
    modelos = [Workstation, EstacaoAllInOne, Notebook, Monitor, Impressora, Periferico, Tablet, Modem, Firewall,
               Roteador, Switch, AccessPoint, Servidor, Projetor, Televisao, NoBreak, Rack, Software]
    resultado = {'Total-de-Ativos': sum(modelo.query.count() for modelo in modelos)}
    for emp in Empresa.query.order_by(Empresa.nome_fantasia.asc()).all():
        bloco = contagens.get(emp.id, {})
        resultado[emp.nome_fantasia] = {
            grupo: {chave: bloco.get(grupo, {}).get(chave, 0) for chave in chaves}
            for grupo, chaves in ((grupo, list(valores.keys()))
                                  for grupo, valores in modelo_vazio.items())
        }
    return jsonify(resultado)


@app.route('/ativos')
def ativos():
    """Webservice com Bearer Token (Chaves de Acesso): lista completa dos ativos
    por empresa em JSON, com todos os campos cadastrados de cada ativo."""
    negado = exigir_bearer()
    if negado:
        return negado
    def serializar(modelo, item):
        registro = {}
        for coluna in modelo.__table__.columns:
            valor = getattr(item, coluna.name)
            if isinstance(valor, (datetime, date)):
                valor = valor.isoformat()
            registro[coluna.name] = valor
        return registro

    def lista(modelo, filtro=None):
        consulta = modelo.query.order_by(modelo.id)
        if filtro is not None:
            consulta = consulta.filter(modelo.tipo_licenciamento == filtro)
        por_empresa = {}
        for item in consulta.all():
            if item.empresa_id is None:
                continue
            por_empresa.setdefault(item.empresa_id, []).append(serializar(modelo, item))
        return por_empresa

    grupos = {}

    def add(modelo, grupo, chave, filtro=None):
        for empresa_id, itens in lista(modelo, filtro).items():
            bloco = grupos.setdefault(empresa_id, {}).setdefault(grupo, {})
            bloco.setdefault(chave, []).extend(itens)

    add(Workstation, 'Hardware', 'Estacoes-de-Trabalho')
    add(EstacaoAllInOne, 'Hardware', 'Estacoes-All-in-One')
    add(Notebook, 'Hardware', 'Notebooks')
    add(Monitor, 'Hardware', 'Monitores')
    add(Impressora, 'Hardware', 'Impressoras')
    add(Periferico, 'Hardware', 'Perifericos')
    add(Tablet, 'Hardware', 'Tablets')
    add(Modem, 'Infraestrutura', 'Modens')
    add(Firewall, 'Infraestrutura', 'Firewalls')
    add(Roteador, 'Infraestrutura', 'Roteadores')
    add(Switch, 'Infraestrutura', 'Switches')
    add(AccessPoint, 'Infraestrutura', 'Access-Points')
    add(Servidor, 'Infraestrutura', 'Servidores')
    add(Projetor, 'Infraestrutura', 'Projetores')
    add(Televisao, 'Infraestrutura', 'Televisoes')
    add(NoBreak, 'Infraestrutura', 'No-Breaks')
    add(Rack, 'Infraestrutura', 'Racks')
    add(Software, 'Software', 'Softwares-Licenciamento-Perpetuo', filtro='Perpétuo')
    add(Software, 'Software', 'Softwares-Por-Assinatura', filtro='Por assinatura')

    modelo_vazio = {'Hardware': {'Estacoes-de-Trabalho': [], 'Estacoes-All-in-One': [], 'Notebooks': [],
                                 'Monitores': [], 'Impressoras': [], 'Perifericos': [], 'Tablets': []},
                    'Infraestrutura': {'Modens': [], 'Firewalls': [], 'Roteadores': [], 'Switches': [],
                                       'Access-Points': [], 'Servidores': [], 'Projetores': [],
                                       'Televisoes': [], 'No-Breaks': [], 'Racks': []},
                    'Software': {'Softwares-Licenciamento-Perpetuo': [],
                                 'Softwares-Por-Assinatura': []}}
    modelos = [Workstation, EstacaoAllInOne, Notebook, Monitor, Impressora, Periferico, Tablet, Modem, Firewall,
               Roteador, Switch, AccessPoint, Servidor, Projetor, Televisao, NoBreak, Rack, Software]
    resultado = {'Total-de-Ativos': sum(modelo.query.count() for modelo in modelos)}
    for emp in Empresa.query.order_by(Empresa.nome_fantasia.asc()).all():
        bloco = grupos.get(emp.id, {})
        resultado[emp.nome_fantasia] = {
            grupo: {chave: bloco.get(grupo, {}).get(chave, []) for chave in chaves}
            for grupo, chaves in ((grupo, list(valores.keys()))
                                  for grupo, valores in modelo_vazio.items())
        }
    return jsonify(resultado)


TABELAS_INFORMACOES = {
    'inf_disco_tamanho': InfDiscoTamanho,
    'inf_hacks': InfHack,
    'inf_impressoras': InfImpressora,
    'inf_memorias_ram_qtde': InfMemoriaRamQtde,
    'inf_memorias_ram_tipo': InfMemoriaRamTipo,
    'inf_nobreaks': InfNoBreak,
    'inf_perifericos': InfPeriferico,
    'inf_perifericos_computador': InfPerifericoComputador,
    'inf_placa_mae_desktop': InfPlacaMaeDesktop,
    'inf_placa_mae_notebook': InfPlacaMaeNotebook,
    'inf_placa_mae_servidor': InfPlacaMaeServidor,
    'inf_placa_rede': InfPlacaRede,
    'inf_placas_de_video': InfPlacaDeVideo,
    'inf_processadores_desktop': InfProcessadorDesktop,
    'inf_processadores_notebook': InfProcessadorNotebook,
    'inf_processadores_servidor': InfProcessadorServidor,
    'inf_servidores': InfServidor,
    'inf_so_desktop': InfSoDesktop,
    'inf_so_servidores': InfSoServidor,
    'inf_sofwares_licencas': InfSofwaresLicencas,
    'inf_tablets_marcas': InfTabletsMarca,
    'inf_tamanho_monitor': InfTamanhoMonitor,
    'inf_televisao_marca': InfTelevisaoMarca,
}


RELATORIOS_EXCEL = {
    'estacoes-all-in-one': {
        'aba': 'Estações All-in-One', 'modelo': EstacaoAllInOne,
        'campos': ['id', 'codigo_interno', '@empresa', 'host', 'marca', 'tipo_memoria_ram', 'memoria_ram',
                   'disco1', 'disco2', 'processador', 'placa_de_video', 'placa_mae', 'sistema_operacional',
                   'placa_de_rede', 'placa_de_rede_ethernet_2', 'placa_de_rede_wifi', 'monitor1', 'monitor2',
                   'teclado', 'mouse', 'usuario', 'local_de_uso', '@perifericos', '@softwares'],
    },
    'estacoes-de-trabalho': {
        'aba': 'Estações de Trabalho', 'modelo': Workstation,
        'campos': ['id', 'codigo_interno', '@empresa', 'host', 'marca', 'tipo_memoria_ram', 'memoria_ram',
                   'disco1', 'disco2', 'processador', 'placa_de_video', 'placa_mae', 'sistema_operacional',
                   'placa_de_rede', 'placa_de_rede_ethernet_2', 'placa_de_rede_wifi', 'monitor1', 'monitor2',
                   'teclado', 'mouse', 'usuario', 'local_de_uso', '@perifericos', '@softwares'],
    },
    'notebooks': {
        'aba': 'Notebooks', 'modelo': Notebook,
        'campos': ['id', 'codigo_interno', '@empresa', 'host', 'marca', 'tipo_memoria_ram', 'memoria_ram',
                   'disco1', 'disco2', 'processador', 'placa_de_video', 'placa_mae', 'sistema_operacional',
                   'placa_de_rede', 'placa_de_rede_ethernet_2', 'placa_de_rede_wifi', 'monitor1', 'monitor2',
                   'teclado', 'mouse', 'usuario', 'local_de_uso', '@perifericos', '@softwares'],
    },
    'monitores': {
        'aba': 'Monitores', 'modelo': Monitor,
        'cabecalhos': ['Código', 'Código Interno', 'Empresa', 'Usuário', 'Marca', 'Tamanho', 'Tipo',
                       'Equipamento', 'Local de Uso'],
        'campos': ['id', 'codigo_interno', '@empresa', 'usuario', 'marca', 'tamanho', 'tipo',
                   '@equipamento', 'local_de_uso'],
    },
    'impressoras': {
        'aba': 'Impressoras', 'modelo': Impressora,
        'campos': ['id', 'codigo_interno', '@empresa', 'marca', 'tipo', 'local_de_uso',
                   '@equipamento', 'outras_informacoes'],
    },
    'perifericos': {
        'aba': 'Periféricos', 'modelo': Periferico,
        'campos': ['id', 'codigo_interno', '@empresa', 'descricao', 'perifericos_multimidia', 'marca',
                   'local_de_uso', '@equipamento', 'usuario', 'outras_informacoes'],
    },
    'tablets': {
        'aba': 'Tablets', 'modelo': Tablet,
        'campos': ['id', 'codigo_interno', '@empresa', 'marca', 'nome', 'cor', 'armazenamento', 'memoria_ram',
                   'resolucao_tela', 'conectividade', 'sistema_operacional', 'local_de_uso', 'usuario',
                   'outras_informacoes'],
    },
    'modens': {
        'aba': 'Modens', 'modelo': Modem,
        'campos': ['id', 'codigo_interno', '@empresa', 'marca', 'host', 'operadora', 'velocidade_rede',
                   'local_de_uso'],
    },
    'firewalls': {
        'aba': 'Firewalls', 'modelo': Firewall,
        'campos': ['id', 'codigo_interno', '@empresa', 'host', 'marca', 'portas_externas', 'portas_internas',
                   'velocidade_rede', 'local_de_uso', 'outras_informacoes'],
    },
    'roteadores': {
        'aba': 'Roteadores', 'modelo': Roteador,
        'campos': ['id', 'codigo_interno', '@empresa', 'host', 'marca', 'tipo', 'portas', 'velocidade_rede',
                   'local_de_uso', 'outras_informacoes'],
    },
    'switches': {
        'aba': 'Switches', 'modelo': Switch,
        'campos': ['id', 'codigo_interno', '@empresa', 'host', 'marca', 'tipo', 'portas', 'velocidade_rede',
                   'local_de_uso', 'outras_informacoes'],
    },
    'access-points': {
        'aba': 'Access Points', 'modelo': AccessPoint,
        'campos': ['id', 'codigo_interno', '@empresa', 'host', 'marca', 'tipo', 'portas', 'velocidade_rede',
                   'local_de_uso', 'outras_informacoes'],
    },
    'servidores': {
        'aba': 'Servidores', 'modelo': Servidor,
        'campos': ['id', 'codigo_interno', '@empresa', 'host', 'marca', 'tipo_memoria_ram', 'memoria_ram',
                   'disco1', 'disco2', 'disco3', 'disco4', 'processador', 'placa_de_video', 'placa_mae',
                   'sistema_operacional', 'placa_de_rede1', 'placa_de_rede2', 'placa_de_rede3',
                   'placa_de_rede4', 'placa_de_rede_wifi', 'monitor1', 'teclado', 'mouse', 'local_de_uso'],
    },
    'projetores': {
        'aba': 'Projetores', 'modelo': Projetor,
        'campos': ['id', 'codigo_interno', '@empresa', 'marca', 'tecnologia_imagem', 'tipo', 'resolucao',
                   'conectividade', 'local_de_uso', 'outras_informacoes'],
    },
    'televisoes': {
        'aba': 'Televisões', 'modelo': Televisao,
        'campos': ['id', 'codigo_interno', '@empresa', 'marca', 'tipo_tecnologia_imagem', 'resolucao',
                   'conectividade_video', 'conectividade_rede', 'local_de_uso', 'outras_informacoes'],
    },
    'nobreaks': {
        'aba': 'No-Breaks', 'modelo': NoBreak,
        'campos': ['id', 'codigo_interno', '@empresa', 'marca', 'capacidade_va', 'potencia_watts',
                   'local_de_uso', 'outras_informacoes'],
    },
    'racks': {
        'aba': 'Racks', 'modelo': Rack,
        'campos': ['id', 'codigo_interno', '@empresa', 'marca', 'cor', 'modelo', 'largura', 'profundidade',
                   'altura', 'us', 'peso_produto', 'capacidade_carga', 'fechadura', 'classificacao_ip',
                   'pes_niveladores', 'local_de_uso', 'outras_informacoes'],
    },
    'softwares': {
        'aba': 'Softwares', 'modelo': Software,
        'campos': ['id', 'codigo_interno', '@empresa', 'categoria', 'marca', 'nome', 'tipo_licenciamento',
                   'validade_assinatura', 'chave', 'codigo_equipamento', 'local_de_uso', 'descricao',
                   'outras_informacoes'],
    },
    'empresas': {
        'aba': 'Empresas', 'modelo': Empresa,
        'campos': ['id', 'cnpj', 'razao_social', 'nome_fantasia', 'cep', 'logradouro', 'bairro', 'cidade',
                   'estado', 'telefones', 'email', 'outras_informacoes'],
    },
    'areas-locais': {
        'aba': 'Áreas e Locais', 'modelo': AreaLocal,
        'campos': ['id', 'empresa', 'area', 'local', 'outras_informacoes'],
    },
    'usuarios': {
        'aba': 'Usuários', 'modelo': Usuario,
        'campos': ['id', 'cpf', 'empresa_que_trabalha', 'nome_completo', 'cep', 'logradouro', 'bairro', 'cidade',
                   'estado', 'telefones', 'email', 'outras_informacoes'],
    },
}


CAMPOS_RELATORIO_MULTIVALOR = {
    'perifericos_multimidia', 'conectividade', 'conectividade_video', 'codigo_equipamento',
}


def mapas_relacionamentos_relatorio(mapa_equipamentos_relatorio):
    perifericos_por_host = {}
    for periferico in Periferico.query.order_by(Periferico.descricao.asc(), Periferico.marca.asc()).all():
        descricao = periferico.descricao or ''
        if periferico.marca:
            descricao = f'{descricao} - {periferico.marca}' if descricao else periferico.marca
        for referencia in (periferico.codigo_equipamento or '').split(','):
            referencia = referencia.strip()
            host = mapa_equipamentos_relatorio.get(referencia, referencia)
            if host:
                perifericos_por_host.setdefault(host, set()).add(descricao)

    workstations = {item.id: item.host for item in Workstation.query.all()}
    notebooks = {item.id: item.host for item in Notebook.query.all()}
    estacoes_all_in_one = {item.id: item.host for item in EstacaoAllInOne.query.all()}
    softwares = {item.id: item for item in Software.query.all()}
    softwares_por_host = {}

    def adicionar_software(host, software):
        if not host or not software:
            return
        rotulo = software.nome or ''
        if software.categoria:
            rotulo = f'{rotulo} ({software.categoria})' if rotulo else software.categoria
        softwares_por_host.setdefault(host, set()).add(rotulo)

    for software in softwares.values():
        if (software.codigo_equipamento or '') != 'Não está instalado':
            for host in (software.codigo_equipamento or '').split(','):
                host = host.strip()
                if host:
                    adicionar_software(mapa_equipamentos_relatorio.get(host, host), software)
        if software.workstation_id:
            adicionar_software(workstations.get(software.workstation_id), software)
        if software.notebook_id:
            adicionar_software(notebooks.get(software.notebook_id), software)

    for vinculo in SoftwareEquipamento.query.all():
        software = softwares.get(vinculo.software_id)
        if vinculo.workstation_id:
            adicionar_software(workstations.get(vinculo.workstation_id), software)
        if vinculo.notebook_id:
            adicionar_software(notebooks.get(vinculo.notebook_id), software)
        if vinculo.estacao_all_in_one_id:
            adicionar_software(estacoes_all_in_one.get(vinculo.estacao_all_in_one_id), software)

    return {
        'perifericos': {host: sorted(valores, key=str.lower) for host, valores in perifericos_por_host.items()},
        'softwares': {host: sorted(valores, key=str.lower) for host, valores in softwares_por_host.items()},
    }


def concatenar_valores_relatorio(valores):
    return '. '.join(str(valor).strip() for valor in valores if valor is not None and str(valor).strip())


def valor_celula_relatorio(item, campo, mapas):
    if campo == '@empresa':
        empresa_id = getattr(item, 'empresa_id', None)
        if empresa_id in mapas['empresas']:
            return mapas['empresas'][empresa_id]
        return getattr(item, 'empresa', None) or getattr(item, 'empresa_que_trabalha', '') or ''
    if campo == '@equipamento':
        referencias = (getattr(item, 'codigo_equipamento', None) or '').split(',')
        return concatenar_valores_relatorio(
            mapas['equipamentos'].get(referencia.strip(), referencia.strip()) for referencia in referencias
        )
    if campo == '@perifericos':
        return concatenar_valores_relatorio(mapas['relacionamentos']['perifericos'].get(item.host, []))
    if campo == '@softwares':
        return concatenar_valores_relatorio(mapas['relacionamentos']['softwares'].get(item.host, []))

    valor = getattr(item, campo, None)
    if valor is None:
        return ''
    if isinstance(valor, str):
        if campo in CAMPOS_RELATORIO_MULTIVALOR:
            return concatenar_valores_relatorio(re.split(r'[,\r\n]+', valor))
        if '\n' in valor or '\r' in valor:
            return concatenar_valores_relatorio(valor.splitlines())
    return valor


def criar_aba_monitores_relatorio(workbook, definicao):
    aba = workbook.create_sheet(definicao['aba'])
    modelo = workbook['Impressoras']
    total_colunas = len(definicao['campos']) + 1
    for coluna in range(1, total_colunas + 1):
        coluna_modelo = min(coluna, modelo.max_column)
        for linha in (1, 2, 3):
            origem = modelo.cell(linha, coluna_modelo)
            destino = aba.cell(linha, coluna)
            if origem.has_style:
                destino._style = copy(origem._style)
            if linha == 3 and coluna == 1:
                destino.value = 'Inventário de Ativos de TI'
        largura = modelo.column_dimensions[get_column_letter(coluna_modelo)].width
        if largura:
            aba.column_dimensions[get_column_letter(coluna)].width = largura
    for linha in (1, 3):
        altura = modelo.row_dimensions[linha].height
        if altura:
            aba.row_dimensions[linha].height = altura
    for indice, cabecalho in enumerate(definicao['cabecalhos'], start=2):
        aba.cell(1, indice).value = cabecalho
    return aba


def carregar_modelo_relatorio(caminho):
    modelo_corrigido = BytesIO()
    mesclagem_invalida = b'<mergeCells count="1"><mergeCell ref="A2:A1048576"/></mergeCells>'
    with ZipFile(caminho, 'r') as arquivo_original, ZipFile(modelo_corrigido, 'w') as arquivo_corrigido:
        for item in arquivo_original.infolist():
            conteudo = arquivo_original.read(item.filename)
            if item.filename.startswith('xl/worksheets/') and item.filename.endswith('.xml'):
                conteudo = conteudo.replace(mesclagem_invalida, b'')
            arquivo_corrigido.writestr(item, conteudo)
    modelo_corrigido.seek(0)
    return load_workbook(modelo_corrigido)


def preencher_aba_relatorio(aba, definicao, itens, mapas):
    campos = definicao['campos']
    if 'cabecalhos' in definicao:
        cabecalhos = definicao['cabecalhos']
    else:
        cabecalhos = []
        for coluna in range(2, aba.max_column + 1):
            cabecalho = aba.cell(1, coluna).value
            if cabecalho is None:
                break
            cabecalhos.append(cabecalho)
    if len(cabecalhos) != len(campos):
        raise ValueError(f'Os campos do relatório não correspondem aos cabeçalhos da aba {aba.title}.')

    if len(itens) > 1:
        estilos_linha = [copy(aba.cell(2, coluna)._style) for coluna in range(1, len(campos) + 2)]
        altura_linha = aba.row_dimensions[2].height
        aba.insert_rows(3, amount=len(itens) - 1)
        for linha in range(3, len(itens) + 2):
            for coluna, estilo in enumerate(estilos_linha, start=1):
                aba.cell(linha, coluna)._style = copy(estilo)
            if altura_linha:
                aba.row_dimensions[linha].height = altura_linha

    for indice, item in enumerate(itens, start=2):
        for coluna, campo in enumerate(campos, start=2):
            valor = valor_celula_relatorio(item, campo, mapas)
            celula = aba.cell(indice, coluna)
            celula.value = valor
            if isinstance(valor, str):
                celula.data_type = 's'


def itens_para_relatorio(definicao):
    modelo = definicao['modelo']
    consulta = modelo.query
    empresa = request.args.get('empresa', 'all')
    if empresa != 'all' and hasattr(modelo, 'empresa_id'):
        consulta = consulta.filter(modelo.empresa_id == int(empresa)) if empresa.isdigit() else consulta.filter(modelo.id == -1)
    if modelo is Software:
        categoria = request.args.get('categoria', 'all')
        if categoria != 'all':
            consulta = consulta.filter(Software.categoria == categoria)
    return consulta.order_by(modelo.id.asc()).all()


@app.route('/relatorios/<relatorio>')
def baixar_relatorio_excel(relatorio):
    definicao = RELATORIOS_EXCEL.get(relatorio)
    if definicao is None:
        abort(404)
    caminho_modelo = Path(app.root_path) / 'modelo' / 'relatorio.xlsx'
    workbook = carregar_modelo_relatorio(caminho_modelo)
    if definicao['aba'] in workbook.sheetnames:
        aba = workbook[definicao['aba']]
    elif relatorio == 'monitores':
        aba = criar_aba_monitores_relatorio(workbook, definicao)
    else:
        workbook.close()
        abort(500, description=f'A aba {definicao["aba"]} não existe no modelo de relatório.')

    for outra_aba in list(workbook.worksheets):
        if outra_aba.title != aba.title:
            workbook.remove(outra_aba)

    itens = itens_para_relatorio(definicao)
    campos = definicao['campos']
    precisa_vinculos = any(campo in ('@perifericos', '@softwares') for campo in campos)
    precisa_equipamentos = precisa_vinculos or '@equipamento' in campos
    mapas_equipamentos = mapa_equipamentos() if precisa_equipamentos else {}
    mapas = {
        'empresas': ({empresa.id: empresa.nome_fantasia for empresa in Empresa.query.all()}
                     if '@empresa' in campos else {}),
        'equipamentos': mapas_equipamentos,
        'relacionamentos': (mapas_relacionamentos_relatorio(mapas_equipamentos) if precisa_vinculos else {}),
    }
    preencher_aba_relatorio(aba, definicao, itens, mapas)

    arquivo = BytesIO()
    workbook.save(arquivo)
    workbook.close()
    arquivo.seek(0)
    nome_arquivo = f'{definicao["aba"]} - {date.today().isoformat()}.xlsx'
    return send_file(
        arquivo,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=nome_arquivo,
    )


@app.route('/informacoes', methods=['GET', 'POST', 'PUT'])
def informacoes():
    """Webservice com Bearer Token (Chaves de Acesso):
    - GET: retorna em JSON todas as informações cadastradas nas 23 tabelas de
      referência inf_* (whitelist TABELAS_INFORMACOES), uma lista por tabela,
      ordenada por codigo; tabela vazia = [].
    - POST: recebe campos de formulário tabela_informacoes e nome; insere a
      linha na tabela de referência whitelist; 201 + {"id": codigo} em sucesso,
      400 para tabela fora da lista permitida ou nome nulo/em branco, 422 para
      falha de gravação.
    - PUT: recebe campos de formulário tabela_informacoes, codigo e nome;
      atualiza o registro com o codigo informado na tabela escolhida; 201 +
      {"id": codigo} em sucesso, 400 para tabela fora da lista permitida ou
      codigo/nome nulo/em branco, 404 se o codigo não existir na tabela,
      422 para falha de gravação."""
    negado = exigir_bearer()
    if negado:
        return negado
    if request.method == 'GET':
        resultado = {}
        for nome_tabela in sorted(TABELAS_INFORMACOES):
            modelo = TABELAS_INFORMACOES[nome_tabela]
            resultado[nome_tabela] = [
                {coluna.name: getattr(item, coluna.name) for coluna in modelo.__table__.columns}
                for item in modelo.query.order_by(modelo.codigo).all()
            ]
        return jsonify(resultado)
    tabela = request.form.get('tabela_informacoes', '')
    nome = request.form.get('nome', '')
    if tabela not in TABELAS_INFORMACOES:
        return jsonify({'erro': 'tabela_informacoes inválida: use uma das tabelas permitidas: '
                                + ', '.join(TABELAS_INFORMACOES)}), 400
    modelo = TABELAS_INFORMACOES[tabela]
    if request.method == 'PUT':
        codigo = request.form.get('codigo', '')
        if not codigo.strip():
            return jsonify({'erro': 'codigo inválido: o campo não pode ser nulo nem em branco'}), 400
        if not nome.strip():
            return jsonify({'erro': 'nome inválido: o campo não pode ser nulo nem em branco'}), 400
        if not codigo.strip().isdigit():
            return jsonify({'erro': 'codigo inválido: use o número do registro (codigo)'}), 400
        item = db.session.get(modelo, int(codigo.strip()))
        if item is None:
            return jsonify({'erro': 'registro não encontrado: nenhum item com o codigo informado '
                                    'na tabela escolhida'}), 404
        try:
            item.nome = nome.strip()
            db.session.commit()
        except Exception:
            db.session.rollback()
            return jsonify({'erro': 'Falha ao gravar o registro'}), 422
        return jsonify({'id': item.codigo}), 201
    if not nome.strip():
        return jsonify({'erro': 'nome inválido: o campo não pode ser nulo nem em branco'}), 400
    try:
        item = modelo(nome=nome.strip())
        db.session.add(item)
        db.session.commit()
    except Exception:
        db.session.rollback()
        return jsonify({'erro': 'Falha ao gravar o registro'}), 422
    return jsonify({'id': item.codigo}), 201


# ---------- Configurações de referência ----------

def configuracao_referencia(tabela):
    configuracao = TABELAS_CONFIGURACOES.get(tabela)
    if configuracao is None:
        abort(404)
    return configuracao


@app.route('/configuracoes/<tabela>')
def list_configuracao_referencia(tabela):
    configuracao = configuracao_referencia(tabela)
    modelo = configuracao['modelo']
    items, ordem, direcao = ordenar_lista(
        modelo,
        {'codigo': modelo.codigo, 'nome': modelo.nome},
        padrao='nome'
    )
    return render_template(
        'configuracoes_referencia.html', configuracao=configuracao,
        items=items, ordem=ordem, direcao=direcao,
        valores_em_uso=valores_referencia_em_uso(tabela)
    )


@app.route('/configuracoes/<tabela>/novo', methods=['GET', 'POST'])
def novo_configuracao_referencia(tabela):
    configuracao = configuracao_referencia(tabela)
    if request.method == 'POST':
        nome = request.form.get('nome', '').strip()
        if not nome:
            flash('Informe o nome.', 'error')
        elif len(nome) > 100:
            flash('O nome deve ter no máximo 100 caracteres.', 'error')
        else:
            try:
                db.session.add(configuracao['modelo'](nome=nome))
                db.session.commit()
                flash('Informação cadastrada com sucesso.', 'success')
                return redirect(url_for('list_configuracao_referencia', tabela=tabela))
            except Exception:
                db.session.rollback()
                flash('Não foi possível cadastrar a informação.', 'error')
    return render_template('configuracao_referencia_form.html', configuracao=configuracao, item=None)


@app.route('/configuracoes/<tabela>/editar/<int:codigo>', methods=['GET', 'POST'])
def editar_configuracao_referencia(tabela, codigo):
    configuracao = configuracao_referencia(tabela)
    item = db.session.get(configuracao['modelo'], codigo)
    if item is None:
        abort(404)
    if request.method == 'POST':
        nome = request.form.get('nome', '').strip()
        if not nome:
            flash('Informe o nome.', 'error')
        elif len(nome) > 100:
            flash('O nome deve ter no máximo 100 caracteres.', 'error')
        else:
            try:
                item.nome = nome
                db.session.commit()
                flash('Informação alterada com sucesso.', 'success')
                return redirect(url_for('list_configuracao_referencia', tabela=tabela))
            except Exception:
                db.session.rollback()
                flash('Não foi possível alterar a informação.', 'error')
    return render_template('configuracao_referencia_form.html', configuracao=configuracao, item=item)


@app.route('/configuracoes/<tabela>/excluir/<int:codigo>', methods=['POST'])
def excluir_configuracao_referencia(tabela, codigo):
    configuracao = configuracao_referencia(tabela)
    item = db.session.get(configuracao['modelo'], codigo)
    if item is None:
        abort(404)
    if item.nome.strip() in valores_referencia_em_uso(tabela):
        flash('Não é possível excluir esta informação: ela está em uso.', 'error')
        return redirect(url_for('list_configuracao_referencia', tabela=tabela))
    try:
        db.session.delete(item)
        db.session.commit()
        flash('Informação excluída com sucesso.', 'success')
    except Exception:
        db.session.rollback()
        flash('Não foi possível excluir a informação.', 'error')
    return redirect(url_for('list_configuracao_referencia', tabela=tabela))


# ---------- Workstations ----------

@app.route('/workstations')
def list_workstations():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(Workstation, {
        'codigo_interno': Workstation.codigo_interno,
        'host': Workstation.host,
        'usuario': Workstation.usuario,
        'processador': Workstation.processador,
        'sistema_operacional': Workstation.sistema_operacional,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('workstations.html', items=items, ordem=ordem, direcao=direcao,
                           empresa_filtro=empresa_filtro)


@app.route('/workstations/novo', methods=['GET', 'POST'])
def novo_workstation():
    if request.method == 'POST':
        usuario_id_form = request.form.get('usuario_id', '')
        usr = Usuario.query.get(int(usuario_id_form)) if usuario_id_form.isdigit() else None
        emp = empresa_do_form()
        al = area_local_do_form()
        w = Workstation(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(Workstation, 'EST'),
            marca=request.form['marca'],
            host=request.form['host'],
            tipo_memoria_ram=request.form['tipo_memoria_ram'],
            memoria_ram=request.form['memoria_ram'],
            disco1=request.form['disco1'],
            disco2=request.form['disco2'],
            processador=request.form['processador'],
            placa_de_video=request.form['placa_de_video'],
            placa_de_rede=request.form['placa_de_rede'],
            placa_de_rede_ethernet_2=request.form['placa_de_rede_ethernet_2'],
            placa_de_rede_wifi=request.form['placa_de_rede_wifi'],
            placa_mae=request.form['placa_mae'],
            sistema_operacional=request.form['sistema_operacional'],
            teclado=request.form['teclado'],
            mouse=request.form['mouse'],
            monitor1=request.form['monitor1'],
            monitor2=request.form['monitor2'],
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            usuario=usr.nome_completo if usr else '',
            usuario_id=usr.id if usr else None,
            empresa_id=emp.id if emp else None,
        )
        try:
            db.session.add(w)
            db.session.flush()
            vincular_monitores([w.monitor1, w.monitor2], f'w-{w.id}', al, usr)
            db.session.commit()
            flash('Estação de trabalho cadastrada com sucesso!', 'success')
            return redirect(url_for('list_workstations'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('workstation_form.html', proximo_codigo=gerar_codigo_interno(Workstation, 'EST'),
                           usuarios=Usuario.query.order_by(Usuario.nome_completo.asc()).all(),
                           **opcoes_workstation())


@app.route('/workstations/editar/<int:id>', methods=['GET', 'POST'])
def editar_workstation(id):
    w = Workstation.query.get_or_404(id)
    monitores_antigos = [w.monitor1, w.monitor2]
    if request.method == 'POST':
        w.codigo_interno = request.form['codigo_interno']
        w.marca = request.form['marca']
        w.host = request.form['host']
        w.tipo_memoria_ram = request.form['tipo_memoria_ram']
        w.memoria_ram = request.form['memoria_ram']
        w.disco1 = request.form['disco1']
        w.disco2 = request.form['disco2']
        w.processador = request.form['processador']
        w.placa_de_video = request.form['placa_de_video']
        w.placa_de_rede = request.form['placa_de_rede']
        w.placa_de_rede_ethernet_2 = request.form['placa_de_rede_ethernet_2']
        w.placa_de_rede_wifi = request.form['placa_de_rede_wifi']
        w.placa_mae = request.form['placa_mae']
        w.sistema_operacional = request.form['sistema_operacional']
        w.teclado = request.form['teclado']
        w.mouse = request.form['mouse']
        w.monitor1 = request.form['monitor1']
        w.monitor2 = request.form['monitor2']
        al = area_local_do_form()
        w.local_de_uso = f'{al.area} - {al.local}' if al else ''
        w.area_local_id = al.id if al else None
        usuario_id_form = request.form.get('usuario_id', '')
        usr = Usuario.query.get(int(usuario_id_form)) if usuario_id_form.isdigit() else None
        emp = empresa_do_form()
        w.usuario = usr.nome_completo if usr else ''
        w.usuario_id = usr.id if usr else None
        w.empresa_id = emp.id if emp else None
        try:
            vincular_monitores([w.monitor1, w.monitor2], f'w-{w.id}', al, usr, monitores_antigos)
            db.session.commit()
            flash('Estação de trabalho atualizada com sucesso!', 'success')
            return redirect(url_for('list_workstations'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('workstation_form.html', item=w,
                           usuarios=Usuario.query.order_by(Usuario.nome_completo.asc()).all(),
                           perifericos_vinculados=itens_vinculados_por_host(w.host, Periferico, 'descricao'),
                           softwares_vinculados=itens_vinculados_por_host(w.host, Software, 'nome'),
                           **opcoes_workstation())


@app.route('/workstations/excluir/<int:id>', methods=['POST'])
def excluir_workstation(id):
    w = Workstation.query.get_or_404(id)
    Software.query.filter_by(workstation_id=w.id).update({'workstation_id': None})
    SoftwareEquipamento.query.filter_by(workstation_id=w.id).delete()
    db.session.delete(w)
    db.session.commit()
    flash('Estação de trabalho excluída com sucesso!', 'success')
    return redirect(url_for('list_workstations'))


# ---------- Estações All-in-One ----------

@app.route('/estacoes-all-in-one')
def list_estacoes_all_in_one():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(EstacaoAllInOne, {
        'codigo_interno': EstacaoAllInOne.codigo_interno,
        'host': EstacaoAllInOne.host,
        'usuario': EstacaoAllInOne.usuario,
        'processador': EstacaoAllInOne.processador,
        'sistema_operacional': EstacaoAllInOne.sistema_operacional,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('estacoes_all_in_one.html', items=items, ordem=ordem, direcao=direcao,
                           empresa_filtro=empresa_filtro)


@app.route('/estacoes-all-in-one/novo', methods=['GET', 'POST'])
def nova_estacao_all_in_one():
    if request.method == 'POST':
        usuario_id_form = request.form.get('usuario_id', '')
        usr = Usuario.query.get(int(usuario_id_form)) if usuario_id_form.isdigit() else None
        emp = empresa_do_form()
        al = area_local_do_form()
        aio = EstacaoAllInOne(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(EstacaoAllInOne, 'AIO'),
            marca=request.form['marca'],
            host=request.form['host'],
            tipo_memoria_ram=request.form['tipo_memoria_ram'],
            memoria_ram=request.form['memoria_ram'],
            disco1=request.form['disco1'],
            disco2=request.form['disco2'],
            processador=request.form['processador'],
            placa_de_video=request.form['placa_de_video'],
            placa_de_rede=request.form['placa_de_rede'],
            placa_de_rede_ethernet_2=request.form['placa_de_rede_ethernet_2'],
            placa_de_rede_wifi=request.form['placa_de_rede_wifi'],
            placa_mae=request.form['placa_mae'],
            sistema_operacional=request.form['sistema_operacional'],
            teclado=request.form['teclado'],
            mouse=request.form['mouse'],
            monitor1=request.form['monitor1'],
            monitor2=request.form['monitor2'],
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            usuario=usr.nome_completo if usr else '',
            usuario_id=usr.id if usr else None,
            empresa_id=emp.id if emp else None,
        )
        try:
            db.session.add(aio)
            db.session.flush()
            # monitor1 é "Resolução do Monitor" (valor literal, sem vínculo); só o Monitor Auxiliar liga monitor.
            vincular_monitores(['', aio.monitor2], f'a-{aio.id}', al, usr)
            db.session.commit()
            flash('Estação All-in-One cadastrada com sucesso!', 'success')
            return redirect(url_for('list_estacoes_all_in_one'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('estacao_all_in_one_form.html', proximo_codigo=gerar_codigo_interno(EstacaoAllInOne, 'AIO'),
                           usuarios=Usuario.query.order_by(Usuario.nome_completo.asc()).all(),
                           **opcoes_estacao_all_in_one())


@app.route('/estacoes-all-in-one/editar/<int:id>', methods=['GET', 'POST'])
def editar_estacao_all_in_one(id):
    aio = EstacaoAllInOne.query.get_or_404(id)
    monitores_antigos = [aio.monitor1, aio.monitor2]
    if request.method == 'POST':
        aio.codigo_interno = request.form['codigo_interno']
        aio.marca = request.form['marca']
        aio.host = request.form['host']
        aio.tipo_memoria_ram = request.form['tipo_memoria_ram']
        aio.memoria_ram = request.form['memoria_ram']
        aio.disco1 = request.form['disco1']
        aio.disco2 = request.form['disco2']
        aio.processador = request.form['processador']
        aio.placa_de_video = request.form['placa_de_video']
        aio.placa_de_rede = request.form['placa_de_rede']
        aio.placa_de_rede_ethernet_2 = request.form['placa_de_rede_ethernet_2']
        aio.placa_de_rede_wifi = request.form['placa_de_rede_wifi']
        aio.placa_mae = request.form['placa_mae']
        aio.sistema_operacional = request.form['sistema_operacional']
        aio.teclado = request.form['teclado']
        aio.mouse = request.form['mouse']
        aio.monitor1 = request.form['monitor1']
        aio.monitor2 = request.form['monitor2']
        al = area_local_do_form()
        aio.local_de_uso = f'{al.area} - {al.local}' if al else ''
        aio.area_local_id = al.id if al else None
        usuario_id_form = request.form.get('usuario_id', '')
        usr = Usuario.query.get(int(usuario_id_form)) if usuario_id_form.isdigit() else None
        emp = empresa_do_form()
        aio.usuario = usr.nome_completo if usr else ''
        aio.usuario_id = usr.id if usr else None
        aio.empresa_id = emp.id if emp else None
        try:
            # monitor1 virou "Resolução do Monitor" (literal): o slot vazio desvincula,
            # na edição, um monitor antigo que tenha sido ligado por aquele campo.
            vincular_monitores(['', aio.monitor2], f'a-{aio.id}', al, usr, monitores_antigos)
            db.session.commit()
            flash('Estação All-in-One atualizada com sucesso!', 'success')
            return redirect(url_for('list_estacoes_all_in_one'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('estacao_all_in_one_form.html', item=aio,
                           usuarios=Usuario.query.order_by(Usuario.nome_completo.asc()).all(),
                           perifericos_vinculados=itens_vinculados_por_host(aio.host, Periferico, 'descricao'),
                           softwares_vinculados=itens_vinculados_por_host(aio.host, Software, 'nome'),
                           **opcoes_estacao_all_in_one())


@app.route('/estacoes-all-in-one/excluir/<int:id>', methods=['POST'])
def excluir_estacao_all_in_one(id):
    aio = EstacaoAllInOne.query.get_or_404(id)
    SoftwareEquipamento.query.filter_by(estacao_all_in_one_id=aio.id).delete()
    db.session.delete(aio)
    db.session.commit()
    flash('Estação All-in-One excluída com sucesso!', 'success')
    return redirect(url_for('list_estacoes_all_in_one'))


# ---------- Notebooks ----------

@app.route('/notebooks')
def list_notebooks():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(Notebook, {
        'codigo_interno': Notebook.codigo_interno,
        'host': Notebook.host,
        'usuario': Notebook.usuario,
        'processador': Notebook.processador,
        'sistema_operacional': Notebook.sistema_operacional,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('notebooks.html', items=items, ordem=ordem, direcao=direcao,
                           empresa_filtro=empresa_filtro)


@app.route('/notebooks/novo', methods=['GET', 'POST'])
def novo_notebook():
    if request.method == 'POST':
        usuario_id_form = request.form.get('usuario_id', '')
        usr = Usuario.query.get(int(usuario_id_form)) if usuario_id_form.isdigit() else None
        emp = empresa_do_form()
        al = area_local_do_form()
        n = Notebook(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(Notebook, 'NBK'),
            marca=request.form['marca'],
            host=request.form['host'],
            tipo_memoria_ram=request.form['tipo_memoria_ram'],
            memoria_ram=request.form['memoria_ram'],
            disco1=request.form['disco1'],
            disco2=request.form['disco2'],
            processador=request.form['processador'],
            placa_de_video=request.form['placa_de_video'],
            placa_de_rede=request.form['placa_de_rede'],
            placa_de_rede_ethernet_2=request.form['placa_de_rede_ethernet_2'],
            placa_de_rede_wifi=request.form['placa_de_rede_wifi'],
            placa_mae=request.form['placa_mae'],
            sistema_operacional=request.form['sistema_operacional'],
            teclado=request.form['teclado'],
            mouse=request.form['mouse'],
            monitor1=request.form['monitor1'],
            monitor2=request.form['monitor2'],
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            usuario=usr.nome_completo if usr else '',
            usuario_id=usr.id if usr else None,
            empresa_id=emp.id if emp else None,
        )
        try:
            db.session.add(n)
            db.session.flush()
            vincular_monitores([n.monitor2], f'n-{n.id}', al, usr)
            db.session.commit()
            flash('Notebook cadastrado com sucesso!', 'success')
            return redirect(url_for('list_notebooks'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('notebook_form.html', proximo_codigo=gerar_codigo_interno(Notebook, 'NBK'),
                           usuarios=Usuario.query.order_by(Usuario.nome_completo.asc()).all(),
                           **opcoes_notebook())


@app.route('/notebooks/editar/<int:id>', methods=['GET', 'POST'])
def editar_notebook(id):
    n = Notebook.query.get_or_404(id)
    monitores_antigos = [n.monitor2]
    if request.method == 'POST':
        n.codigo_interno = request.form['codigo_interno']
        n.marca = request.form['marca']
        n.host = request.form['host']
        n.tipo_memoria_ram = request.form['tipo_memoria_ram']
        n.memoria_ram = request.form['memoria_ram']
        n.disco1 = request.form['disco1']
        n.disco2 = request.form['disco2']
        n.processador = request.form['processador']
        n.placa_de_video = request.form['placa_de_video']
        n.placa_de_rede = request.form['placa_de_rede']
        n.placa_de_rede_ethernet_2 = request.form['placa_de_rede_ethernet_2']
        n.placa_de_rede_wifi = request.form['placa_de_rede_wifi']
        n.placa_mae = request.form['placa_mae']
        n.sistema_operacional = request.form['sistema_operacional']
        n.teclado = request.form['teclado']
        n.mouse = request.form['mouse']
        n.monitor1 = request.form['monitor1']
        n.monitor2 = request.form['monitor2']
        al = area_local_do_form()
        n.local_de_uso = f'{al.area} - {al.local}' if al else ''
        n.area_local_id = al.id if al else None
        usuario_id_form = request.form.get('usuario_id', '')
        usr = Usuario.query.get(int(usuario_id_form)) if usuario_id_form.isdigit() else None
        emp = empresa_do_form()
        n.usuario = usr.nome_completo if usr else ''
        n.usuario_id = usr.id if usr else None
        n.empresa_id = emp.id if emp else None
        try:
            vincular_monitores([n.monitor2], f'n-{n.id}', al, usr, monitores_antigos)
            db.session.commit()
            flash('Notebook atualizado com sucesso!', 'success')
            return redirect(url_for('list_notebooks'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('notebook_form.html', item=n,
                           usuarios=Usuario.query.order_by(Usuario.nome_completo.asc()).all(),
                           perifericos_vinculados=itens_vinculados_por_host(n.host, Periferico, 'descricao'),
                           softwares_vinculados=itens_vinculados_por_host(n.host, Software, 'nome'),
                           **opcoes_notebook())


@app.route('/notebooks/excluir/<int:id>', methods=['POST'])
def excluir_notebook(id):
    n = Notebook.query.get_or_404(id)
    Software.query.filter_by(notebook_id=n.id).update({'notebook_id': None})
    SoftwareEquipamento.query.filter_by(notebook_id=n.id).delete()
    db.session.delete(n)
    db.session.commit()
    flash('Notebook excluído com sucesso!', 'success')
    return redirect(url_for('list_notebooks'))


# ---------- Monitores ----------

@app.route('/monitores')
def list_monitores():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(Monitor, {
        'codigo_interno': Monitor.codigo_interno,
        'usuario': Monitor.usuario,
        'tamanho': Monitor.tamanho,
        'tipo': Monitor.tipo,
        'codigo_equipamento': Monitor.codigo_equipamento,
        'local_de_uso': Monitor.local_de_uso,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('monitores.html', items=items, ordem=ordem, direcao=direcao,
                           equipamentos=mapa_equipamentos(), empresa_filtro=empresa_filtro)


@app.route('/monitores/novo', methods=['GET', 'POST'])
def novo_monitor():
    if request.method == 'POST':
        usuario_id_form = request.form.get('usuario_id', '')
        usr = Usuario.query.get(int(usuario_id_form)) if usuario_id_form.isdigit() else None
        emp = empresa_do_form()
        al = area_local_do_form()
        m = Monitor(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(Monitor, 'MON'),
            marca=request.form['marca'],
            tamanho=request.form['tamanho'],
            tipo=request.form['tipo'],
            codigo_equipamento=request.form.get('codigo_equipamento', ''),
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            usuario=usr.nome_completo if usr else '',
            usuario_id=usr.id if usr else None,
            empresa_id=emp.id if emp else None,
        )
        try:
            db.session.add(m)
            db.session.commit()
            flash('Monitor cadastrado com sucesso!', 'success')
            return redirect(url_for('list_monitores'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('monitor_form.html', proximo_codigo=gerar_codigo_interno(Monitor, 'MON'),
                           usuarios=Usuario.query.order_by(Usuario.nome_completo.asc()).all(),
                           **opcoes_monitor())


@app.route('/monitores/editar/<int:id>', methods=['GET', 'POST'])
def editar_monitor(id):
    m = Monitor.query.get_or_404(id)
    if request.method == 'POST':
        m.codigo_interno = request.form['codigo_interno']
        m.marca = request.form['marca']
        m.tamanho = request.form['tamanho']
        m.tipo = request.form['tipo']
        m.codigo_equipamento = request.form.get('codigo_equipamento', '')
        al = area_local_do_form()
        m.local_de_uso = f'{al.area} - {al.local}' if al else ''
        m.area_local_id = al.id if al else None
        usuario_id_form = request.form.get('usuario_id', '')
        usr = Usuario.query.get(int(usuario_id_form)) if usuario_id_form.isdigit() else None
        emp = empresa_do_form()
        m.usuario = usr.nome_completo if usr else ''
        m.usuario_id = usr.id if usr else None
        m.empresa_id = emp.id if emp else None
        try:
            db.session.commit()
            flash('Monitor atualizado com sucesso!', 'success')
            return redirect(url_for('list_monitores'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('monitor_form.html', item=m,
                           usuarios=Usuario.query.order_by(Usuario.nome_completo.asc()).all(),
                           **opcoes_monitor())


@app.route('/monitores/excluir/<int:id>', methods=['POST'])
def excluir_monitor(id):
    m = Monitor.query.get_or_404(id)
    db.session.delete(m)
    db.session.commit()
    flash('Monitor excluído com sucesso!', 'success')
    return redirect(url_for('list_monitores'))


# ---------- Periféricos ----------

@app.route('/perifericos')
def list_perifericos():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(Periferico, {
        'codigo_interno': Periferico.codigo_interno,
        'descricao': Periferico.descricao,
        'perifericos_multimidia': Periferico.perifericos_multimidia,
        'codigo_equipamento': Periferico.codigo_equipamento,
        'usuario': Periferico.usuario,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('perifericos.html', items=items, ordem=ordem, direcao=direcao,
                           equipamentos=mapa_equipamentos(), empresa_filtro=empresa_filtro)


@app.route('/perifericos/novo', methods=['GET', 'POST'])
def novo_periferico():
    if request.method == 'POST':
        perifericos = ', '.join(request.form.getlist('perifericos_multimidia'))
        usuario_id_form = request.form.get('usuario_id', '')
        usr = Usuario.query.get(int(usuario_id_form)) if usuario_id_form.isdigit() else None
        emp = empresa_do_form()
        al = area_local_do_form()
        p = Periferico(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(Periferico, 'PER'),
            descricao=request.form['descricao'],
            marca=request.form['marca'],
            perifericos_multimidia=perifericos,
            codigo_equipamento=request.form.get('codigo_equipamento', ''),
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            usuario=usr.nome_completo if usr else '',
            usuario_id=usr.id if usr else None,
            empresa_id=emp.id if emp else None,
            outras_informacoes=request.form.get('outras_informacoes', ''),
        )
        try:
            db.session.add(p)
            db.session.commit()
            flash('Periférico cadastrado com sucesso!', 'success')
            return redirect(url_for('list_perifericos'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('periferico_form.html', proximo_codigo=gerar_codigo_interno(Periferico, 'PER'),
                           usuarios=Usuario.query.order_by(Usuario.nome_completo.asc()).all(),
                           **opcoes_periferico())


@app.route('/perifericos/editar/<int:id>', methods=['GET', 'POST'])
def editar_periferico(id):
    p = Periferico.query.get_or_404(id)
    if request.method == 'POST':
        perifericos = ', '.join(request.form.getlist('perifericos_multimidia'))
        p.codigo_interno = request.form['codigo_interno']
        p.descricao = request.form['descricao']
        p.marca = request.form['marca']
        p.perifericos_multimidia = perifericos
        p.codigo_equipamento = request.form.get('codigo_equipamento', '')
        al = area_local_do_form()
        p.local_de_uso = f'{al.area} - {al.local}' if al else ''
        p.area_local_id = al.id if al else None
        usuario_id_form = request.form.get('usuario_id', '')
        usr = Usuario.query.get(int(usuario_id_form)) if usuario_id_form.isdigit() else None
        emp = empresa_do_form()
        p.usuario = usr.nome_completo if usr else ''
        p.usuario_id = usr.id if usr else None
        p.empresa_id = emp.id if emp else None
        p.outras_informacoes = request.form.get('outras_informacoes', '')
        try:
            db.session.commit()
            flash('Periférico atualizado com sucesso!', 'success')
            return redirect(url_for('list_perifericos'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('periferico_form.html', item=p,
                           usuarios=Usuario.query.order_by(Usuario.nome_completo.asc()).all(),
                           **opcoes_periferico())


@app.route('/perifericos/excluir/<int:id>', methods=['POST'])
def excluir_periferico(id):
    p = Periferico.query.get_or_404(id)
    db.session.delete(p)
    db.session.commit()
    flash('Periférico excluído com sucesso!', 'success')
    return redirect(url_for('list_perifericos'))


# ---------- Impressoras ----------

@app.route('/impressoras')
def list_impressoras():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(Impressora, {
        'codigo_interno': Impressora.codigo_interno,
        'marca': Impressora.marca,
        'tipo': Impressora.tipo,
        'codigo_equipamento': Impressora.codigo_equipamento,
        'local_de_uso': Impressora.local_de_uso,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('impressoras.html', items=items, ordem=ordem, direcao=direcao,
                           equipamentos=mapa_equipamentos(), empresa_filtro=empresa_filtro)


@app.route('/impressoras/novo', methods=['GET', 'POST'])
def nova_impressora():
    if request.method == 'POST':
        emp = empresa_do_form()
        al = area_local_do_form()
        i = Impressora(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(Impressora, 'IMP'),
            marca=request.form['marca'],
            tipo=request.form['tipo'],
            codigo_equipamento=request.form.get('codigo_equipamento', ''),
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            empresa_id=emp.id if emp else None,
            outras_informacoes=request.form.get('outras_informacoes', ''),
        )
        try:
            db.session.add(i)
            db.session.commit()
            flash('Impressora cadastrada com sucesso!', 'success')
            return redirect(url_for('list_impressoras'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('impressora_form.html', proximo_codigo=gerar_codigo_interno(Impressora, 'IMP'),
                           **opcoes_impressora())


@app.route('/impressoras/editar/<int:id>', methods=['GET', 'POST'])
def editar_impressora(id):
    i = Impressora.query.get_or_404(id)
    if request.method == 'POST':
        i.codigo_interno = request.form['codigo_interno']
        i.marca = request.form['marca']
        i.tipo = request.form['tipo']
        i.codigo_equipamento = request.form.get('codigo_equipamento', '')
        emp = empresa_do_form()
        al = area_local_do_form()
        i.local_de_uso = f'{al.area} - {al.local}' if al else ''
        i.area_local_id = al.id if al else None
        i.empresa_id = emp.id if emp else None
        i.outras_informacoes = request.form.get('outras_informacoes', '')
        try:
            db.session.commit()
            flash('Impressora atualizada com sucesso!', 'success')
            return redirect(url_for('list_impressoras'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('impressora_form.html', item=i, **opcoes_impressora())


@app.route('/impressoras/excluir/<int:id>', methods=['POST'])
def excluir_impressora(id):
    i = Impressora.query.get_or_404(id)
    db.session.delete(i)
    db.session.commit()
    flash('Impressora excluída com sucesso!', 'success')
    return redirect(url_for('list_impressoras'))


# ---------- Tablets ----------

@app.route('/tablets')
def list_tablets():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(Tablet, {
        'codigo_interno': Tablet.codigo_interno,
        'usuario': Tablet.usuario,
        'nome': Tablet.nome,
        'marca': Tablet.marca,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('tablets.html', items=items, ordem=ordem, direcao=direcao,
                           empresa_filtro=empresa_filtro)


@app.route('/tablets/novo', methods=['GET', 'POST'])
def nova_tablet():
    if request.method == 'POST':
        conectividade = ', '.join(v for v in request.form.getlist('conectividade') if v)
        sistema_operacional = request.form.get('sistema_operacional', '')
        if not sistema_operacional:
            flash('Selecione o Sistema Operacional.', 'error')
        elif not conectividade:
            flash('Selecione ao menos uma opção de Conectividade.', 'error')
        else:
            usuario_id_form = request.form.get('usuario_id', '')
            usr = Usuario.query.get(int(usuario_id_form)) if usuario_id_form.isdigit() else None
            emp = empresa_do_form()
            al = area_local_do_form()
            tb = Tablet(
                codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(Tablet, 'TBL'),
                empresa_id=emp.id if emp else None,
                marca=request.form['marca'],
                nome=request.form['nome'],
                cor=request.form['cor'],
                armazenamento=request.form['armazenamento'],
                memoria_ram=request.form['memoria_ram'],
                resolucao_tela=request.form['resolucao_tela'],
                conectividade=conectividade,
                sistema_operacional=sistema_operacional,
                local_de_uso=f'{al.area} - {al.local}' if al else '',
                area_local_id=al.id if al else None,
                usuario=usr.nome_completo if usr else '',
                usuario_id=usr.id if usr else None,
                outras_informacoes=request.form.get('outras_informacoes', ''),
            )
            try:
                db.session.add(tb)
                db.session.commit()
                flash('Tablet cadastrado com sucesso!', 'success')
                return redirect(url_for('list_tablets'))
            except Exception as e:
                db.session.rollback()
                flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('tablet_form.html', proximo_codigo=gerar_codigo_interno(Tablet, 'TBL'),
                           **opcoes_tablet())


@app.route('/tablets/editar/<int:id>', methods=['GET', 'POST'])
def editar_tablet(id):
    tb = Tablet.query.get_or_404(id)
    if request.method == 'POST':
        conectividade = ', '.join(v for v in request.form.getlist('conectividade') if v)
        sistema_operacional = request.form.get('sistema_operacional', '')
        if not sistema_operacional:
            flash('Selecione o Sistema Operacional.', 'error')
        elif not conectividade:
            flash('Selecione ao menos uma opção de Conectividade.', 'error')
        else:
            tb.codigo_interno = request.form['codigo_interno']
            emp = empresa_do_form()
            tb.empresa_id = emp.id if emp else None
            tb.marca = request.form['marca']
            tb.nome = request.form['nome']
            tb.cor = request.form['cor']
            tb.armazenamento = request.form['armazenamento']
            tb.memoria_ram = request.form['memoria_ram']
            tb.resolucao_tela = request.form['resolucao_tela']
            tb.conectividade = conectividade
            tb.sistema_operacional = sistema_operacional
            usuario_id_form = request.form.get('usuario_id', '')
            usr = Usuario.query.get(int(usuario_id_form)) if usuario_id_form.isdigit() else None
            al = area_local_do_form()
            tb.local_de_uso = f'{al.area} - {al.local}' if al else ''
            tb.area_local_id = al.id if al else None
            tb.usuario = usr.nome_completo if usr else ''
            tb.usuario_id = usr.id if usr else None
            tb.outras_informacoes = request.form.get('outras_informacoes', '')
            try:
                db.session.commit()
                flash('Tablet atualizado com sucesso!', 'success')
                return redirect(url_for('list_tablets'))
            except Exception as e:
                db.session.rollback()
                flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('tablet_form.html', item=tb, **opcoes_tablet())


@app.route('/tablets/excluir/<int:id>', methods=['POST'])
def excluir_tablet(id):
    tb = Tablet.query.get_or_404(id)
    db.session.delete(tb)
    db.session.commit()
    flash('Tablet excluído com sucesso!', 'success')
    return redirect(url_for('list_tablets'))


# ---------- Servidores ----------

@app.route('/servidores')
def list_servidores():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(Servidor, {
        'codigo_interno': Servidor.codigo_interno,
        'host': Servidor.host,
        'sistema_operacional': Servidor.sistema_operacional,
        'local_de_uso': Servidor.local_de_uso,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('servidores.html', items=items, ordem=ordem, direcao=direcao,
                           empresa_filtro=empresa_filtro)


@app.route('/servidores/novo', methods=['GET', 'POST'])
def novo_servidor():
    if request.method == 'POST':
        emp = empresa_do_form()
        al = area_local_do_form()
        s = Servidor(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(Servidor, 'SRV'),
            marca=request.form['marca'],
            host=request.form['host'],
            tipo_memoria_ram=request.form['tipo_memoria_ram'],
            memoria_ram=request.form['memoria_ram'],
            disco1=request.form['disco1'],
            disco2=request.form['disco2'],
            disco3=request.form['disco3'],
            disco4=request.form['disco4'],
            processador=request.form['processador'],
            placa_de_video=request.form['placa_de_video'],
            placa_de_rede1=request.form['placa_de_rede1'],
            placa_de_rede2=request.form['placa_de_rede2'],
            placa_de_rede3=request.form['placa_de_rede3'],
            placa_de_rede4=request.form['placa_de_rede4'],
            placa_de_rede_wifi=request.form['placa_de_rede_wifi'],
            placa_mae=request.form['placa_mae'],
            sistema_operacional=request.form['sistema_operacional'],
            teclado=request.form['teclado'],
            mouse=request.form['mouse'],
            monitor1=request.form['monitor1'],
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            empresa_id=emp.id if emp else None,
        )
        try:
            db.session.add(s)
            db.session.flush()
            vincular_monitores([s.monitor1], f's-{s.id}', al, None)
            db.session.commit()
            flash('Servidor cadastrado com sucesso!', 'success')
            return redirect(url_for('list_servidores'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('servidor_form.html', proximo_codigo=gerar_codigo_interno(Servidor, 'SRV'),
                           **opcoes_servidor())


@app.route('/servidores/editar/<int:id>', methods=['GET', 'POST'])
def editar_servidor(id):
    s = Servidor.query.get_or_404(id)
    monitores_antigos = [s.monitor1]
    if request.method == 'POST':
        s.codigo_interno = request.form['codigo_interno']
        s.marca = request.form['marca']
        s.host = request.form['host']
        s.tipo_memoria_ram = request.form['tipo_memoria_ram']
        s.memoria_ram = request.form['memoria_ram']
        s.disco1 = request.form['disco1']
        s.disco2 = request.form['disco2']
        s.disco3 = request.form['disco3']
        s.disco4 = request.form['disco4']
        s.processador = request.form['processador']
        s.placa_de_video = request.form['placa_de_video']
        s.placa_de_rede1 = request.form['placa_de_rede1']
        s.placa_de_rede2 = request.form['placa_de_rede2']
        s.placa_de_rede3 = request.form['placa_de_rede3']
        s.placa_de_rede4 = request.form['placa_de_rede4']
        s.placa_de_rede_wifi = request.form['placa_de_rede_wifi']
        s.placa_mae = request.form['placa_mae']
        s.sistema_operacional = request.form['sistema_operacional']
        s.teclado = request.form['teclado']
        s.mouse = request.form['mouse']
        emp = empresa_do_form()
        s.monitor1 = request.form['monitor1']
        al = area_local_do_form()
        s.local_de_uso = f'{al.area} - {al.local}' if al else ''
        s.area_local_id = al.id if al else None
        s.empresa_id = emp.id if emp else None
        try:
            vincular_monitores([s.monitor1], f's-{s.id}', al, None, monitores_antigos)
            db.session.commit()
            flash('Servidor atualizado com sucesso!', 'success')
            return redirect(url_for('list_servidores'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('servidor_form.html', item=s, **opcoes_servidor())


@app.route('/servidores/excluir/<int:id>', methods=['POST'])
def excluir_servidor(id):
    s = Servidor.query.get_or_404(id)
    db.session.delete(s)
    db.session.commit()
    flash('Servidor excluído com sucesso!', 'success')
    return redirect(url_for('list_servidores'))


# ---------- Modens ----------

@app.route('/modens')
def list_modems():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(Modem, {
        'codigo_interno': Modem.codigo_interno,
        'host': Modem.host,
        'operadora': Modem.operadora,
        'velocidade_rede': Modem.velocidade_rede,
        'local_de_uso': Modem.local_de_uso,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('modems.html', items=items, ordem=ordem, direcao=direcao,
                           empresa_filtro=empresa_filtro)


@app.route('/modens/novo', methods=['GET', 'POST'])
def novo_modem():
    if request.method == 'POST':
        emp = empresa_do_form()
        al = area_local_do_form()
        m = Modem(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(Modem, 'MOD'),
            marca=request.form['marca'],
            host=request.form['host'],
            operadora=request.form['operadora'],
            velocidade_rede=request.form['velocidade_rede'],
            empresa_id=emp.id if emp else None,
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
        )
        try:
            db.session.add(m)
            db.session.commit()
            flash('Modem cadastrado com sucesso!', 'success')
            return redirect(url_for('list_modems'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('modem_form.html', proximo_codigo=gerar_codigo_interno(Modem, 'MOD'))


@app.route('/modens/editar/<int:id>', methods=['GET', 'POST'])
def editar_modem(id):
    m = Modem.query.get_or_404(id)
    if request.method == 'POST':
        m.codigo_interno = request.form['codigo_interno']
        m.marca = request.form['marca']
        m.host = request.form['host']
        m.operadora = request.form['operadora']
        m.velocidade_rede = request.form['velocidade_rede']
        emp = empresa_do_form()
        m.empresa_id = emp.id if emp else None
        al = area_local_do_form()
        m.local_de_uso = f'{al.area} - {al.local}' if al else ''
        m.area_local_id = al.id if al else None
        try:
            db.session.commit()
            flash('Modem atualizado com sucesso!', 'success')
            return redirect(url_for('list_modems'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('modem_form.html', item=m)


@app.route('/modens/excluir/<int:id>', methods=['POST'])
def excluir_modem(id):
    m = Modem.query.get_or_404(id)
    db.session.delete(m)
    db.session.commit()
    flash('Modem excluído com sucesso!', 'success')
    return redirect(url_for('list_modems'))


# ---------- Roteadores ----------

@app.route('/roteadores')
def list_roteadores():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(Roteador, {
        'codigo_interno': Roteador.codigo_interno,
        'marca': Roteador.marca,
        'tipo': Roteador.tipo,
        'velocidade_rede': Roteador.velocidade_rede,
        'local_de_uso': Roteador.local_de_uso,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('roteadores.html', items=items, ordem=ordem, direcao=direcao,
                           empresa_filtro=empresa_filtro)


@app.route('/roteadores/novo', methods=['GET', 'POST'])
def novo_roteador():
    if request.method == 'POST':
        emp = empresa_do_form()
        al = area_local_do_form()
        r = Roteador(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(Roteador, 'ROT'),
            host=request.form.get('host', '').strip() or None,
            marca=request.form['marca'],
            tipo=request.form['tipo'],
            portas=request.form['portas'],
            velocidade_rede=request.form['velocidade_rede'],
            empresa_id=emp.id if emp else None,
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            outras_informacoes=request.form.get('outras_informacoes', ''),
        )
        try:
            db.session.add(r)
            db.session.commit()
            flash('Roteador cadastrado com sucesso!', 'success')
            return redirect(url_for('list_roteadores'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('roteador_form.html', proximo_codigo=gerar_codigo_interno(Roteador, 'ROT'))


@app.route('/roteadores/editar/<int:id>', methods=['GET', 'POST'])
def editar_roteador(id):
    r = Roteador.query.get_or_404(id)
    if request.method == 'POST':
        r.codigo_interno = request.form['codigo_interno']
        r.host = request.form.get('host', '').strip() or None
        r.marca = request.form['marca']
        r.tipo = request.form['tipo']
        r.portas = request.form['portas']
        r.velocidade_rede = request.form['velocidade_rede']
        emp = empresa_do_form()
        r.empresa_id = emp.id if emp else None
        al = area_local_do_form()
        r.local_de_uso = f'{al.area} - {al.local}' if al else ''
        r.area_local_id = al.id if al else None
        r.outras_informacoes = request.form.get('outras_informacoes', '')
        try:
            db.session.commit()
            flash('Roteador atualizado com sucesso!', 'success')
            return redirect(url_for('list_roteadores'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('roteador_form.html', item=r)


@app.route('/roteadores/excluir/<int:id>', methods=['POST'])
def excluir_roteador(id):
    r = Roteador.query.get_or_404(id)
    db.session.delete(r)
    db.session.commit()
    flash('Roteador excluído com sucesso!', 'success')
    return redirect(url_for('list_roteadores'))


# ---------- Switches ----------

@app.route('/switches')
def list_switches():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(Switch, {
        'codigo_interno': Switch.codigo_interno,
        'marca': Switch.marca,
        'tipo': Switch.tipo,
        'velocidade_rede': Switch.velocidade_rede,
        'local_de_uso': Switch.local_de_uso,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('switches.html', items=items, ordem=ordem, direcao=direcao,
                           empresa_filtro=empresa_filtro)


@app.route('/switches/novo', methods=['GET', 'POST'])
def novo_switch():
    if request.method == 'POST':
        emp = empresa_do_form()
        al = area_local_do_form()
        s = Switch(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(Switch, 'SWT'),
            host=request.form.get('host', '').strip() or None,
            marca=request.form['marca'],
            tipo=request.form['tipo'],
            portas=request.form['portas'],
            velocidade_rede=request.form['velocidade_rede'],
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            empresa_id=emp.id if emp else None,
            outras_informacoes=request.form.get('outras_informacoes', ''),
        )
        try:
            db.session.add(s)
            db.session.commit()
            flash('Switch cadastrado com sucesso!', 'success')
            return redirect(url_for('list_switches'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('switch_form.html', proximo_codigo=gerar_codigo_interno(Switch, 'SWT'))


@app.route('/switches/editar/<int:id>', methods=['GET', 'POST'])
def editar_switch(id):
    s = Switch.query.get_or_404(id)
    if request.method == 'POST':
        s.codigo_interno = request.form['codigo_interno']
        s.host = request.form.get('host', '').strip() or None
        s.marca = request.form['marca']
        s.tipo = request.form['tipo']
        s.portas = request.form['portas']
        s.velocidade_rede = request.form['velocidade_rede']
        emp = empresa_do_form()
        al = area_local_do_form()
        s.local_de_uso = f'{al.area} - {al.local}' if al else ''
        s.area_local_id = al.id if al else None
        s.empresa_id = emp.id if emp else None
        s.outras_informacoes = request.form.get('outras_informacoes', '')
        try:
            db.session.commit()
            flash('Switch atualizado com sucesso!', 'success')
            return redirect(url_for('list_switches'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('switch_form.html', item=s)


@app.route('/switches/excluir/<int:id>', methods=['POST'])
def excluir_switch(id):
    s = Switch.query.get_or_404(id)
    db.session.delete(s)
    db.session.commit()
    flash('Switch excluído com sucesso!', 'success')
    return redirect(url_for('list_switches'))


# ---------- Firewalls ----------

@app.route('/firewalls')
def list_firewalls():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(Firewall, {
        'codigo_interno': Firewall.codigo_interno,
        'marca': Firewall.marca,
        'velocidade_rede': Firewall.velocidade_rede,
        'local_de_uso': Firewall.local_de_uso,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('firewalls.html', items=items, ordem=ordem, direcao=direcao,
                           empresa_filtro=empresa_filtro)


@app.route('/firewalls/novo', methods=['GET', 'POST'])
def novo_firewall():
    if request.method == 'POST':
        emp = empresa_do_form()
        al = area_local_do_form()
        f = Firewall(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(Firewall, 'FWL'),
            host=request.form.get('host', '').strip() or None,
            marca=request.form['marca'],
            portas_externas=request.form['portas_externas'],
            portas_internas=request.form['portas_internas'],
            velocidade_rede=request.form['velocidade_rede'],
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            empresa_id=emp.id if emp else None,
            outras_informacoes=request.form.get('outras_informacoes', ''),
        )
        try:
            db.session.add(f)
            db.session.commit()
            flash('Firewall cadastrado com sucesso!', 'success')
            return redirect(url_for('list_firewalls'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('firewall_form.html', proximo_codigo=gerar_codigo_interno(Firewall, 'FWL'))


@app.route('/firewalls/editar/<int:id>', methods=['GET', 'POST'])
def editar_firewall(id):
    f = Firewall.query.get_or_404(id)
    if request.method == 'POST':
        f.codigo_interno = request.form['codigo_interno']
        f.host = request.form.get('host', '').strip() or None
        f.marca = request.form['marca']
        f.portas_externas = request.form['portas_externas']
        f.portas_internas = request.form['portas_internas']
        f.velocidade_rede = request.form['velocidade_rede']
        emp = empresa_do_form()
        al = area_local_do_form()
        f.local_de_uso = f'{al.area} - {al.local}' if al else ''
        f.area_local_id = al.id if al else None
        f.empresa_id = emp.id if emp else None
        f.outras_informacoes = request.form.get('outras_informacoes', '')
        try:
            db.session.commit()
            flash('Firewall atualizado com sucesso!', 'success')
            return redirect(url_for('list_firewalls'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('firewall_form.html', item=f)


@app.route('/firewalls/excluir/<int:id>', methods=['POST'])
def excluir_firewall(id):
    f = Firewall.query.get_or_404(id)
    db.session.delete(f)
    db.session.commit()
    flash('Firewall excluído com sucesso!', 'success')
    return redirect(url_for('list_firewalls'))


# ---------- Access Points ----------

@app.route('/access-points')
def list_access_points():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(AccessPoint, {
        'codigo_interno': AccessPoint.codigo_interno,
        'marca': AccessPoint.marca,
        'tipo': AccessPoint.tipo,
        'velocidade_rede': AccessPoint.velocidade_rede,
        'local_de_uso': AccessPoint.local_de_uso,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('access_points.html', items=items, ordem=ordem, direcao=direcao,
                           empresa_filtro=empresa_filtro)


@app.route('/access-points/novo', methods=['GET', 'POST'])
def novo_access_point():
    if request.method == 'POST':
        emp = empresa_do_form()
        al = area_local_do_form()
        ap = AccessPoint(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(AccessPoint, 'ACP'),
            host=request.form.get('host', '').strip() or None,
            marca=request.form['marca'],
            tipo=request.form['tipo'],
            portas=request.form['portas'],
            velocidade_rede=request.form['velocidade_rede'],
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            empresa_id=emp.id if emp else None,
            outras_informacoes=request.form.get('outras_informacoes', ''),
        )
        try:
            db.session.add(ap)
            db.session.commit()
            flash('Access Point cadastrado com sucesso!', 'success')
            return redirect(url_for('list_access_points'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('access_point_form.html', proximo_codigo=gerar_codigo_interno(AccessPoint, 'ACP'))


@app.route('/access-points/editar/<int:id>', methods=['GET', 'POST'])
def editar_access_point(id):
    ap = AccessPoint.query.get_or_404(id)
    if request.method == 'POST':
        ap.codigo_interno = request.form['codigo_interno']
        ap.host = request.form.get('host', '').strip() or None
        ap.marca = request.form['marca']
        ap.tipo = request.form['tipo']
        ap.portas = request.form['portas']
        ap.velocidade_rede = request.form['velocidade_rede']
        emp = empresa_do_form()
        al = area_local_do_form()
        ap.local_de_uso = f'{al.area} - {al.local}' if al else ''
        ap.area_local_id = al.id if al else None
        ap.empresa_id = emp.id if emp else None
        ap.outras_informacoes = request.form.get('outras_informacoes', '')
        try:
            db.session.commit()
            flash('Access Point atualizado com sucesso!', 'success')
            return redirect(url_for('list_access_points'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('access_point_form.html', item=ap)


@app.route('/access-points/excluir/<int:id>', methods=['POST'])
def excluir_access_point(id):
    ap = AccessPoint.query.get_or_404(id)
    db.session.delete(ap)
    db.session.commit()
    flash('Access Point excluído com sucesso!', 'success')
    return redirect(url_for('list_access_points'))


# ---------- No-Breaks ----------

@app.route('/nobreaks')
def list_nobreaks():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(NoBreak, {
        'codigo_interno': NoBreak.codigo_interno,
        'empresa': NoBreak.empresa_id,
        'local_de_uso': NoBreak.local_de_uso,
        'marca': NoBreak.marca,
        'capacidade_va': NoBreak.capacidade_va,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    mapa_empresas = {e.id: e.nome_fantasia for e in Empresa.query.all()}
    return render_template('nobreaks.html', items=items, ordem=ordem, direcao=direcao,
                           mapa_empresas=mapa_empresas, empresa_filtro=empresa_filtro)


@app.route('/nobreaks/novo', methods=['GET', 'POST'])
def novo_nobreak():
    if request.method == 'POST':
        emp = empresa_do_form()
        al = area_local_do_form()
        nb = NoBreak(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(NoBreak, 'NBR'),
            marca=request.form['marca'],
            capacidade_va=request.form['capacidade_va'],
            potencia_watts=request.form['potencia_watts'],
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            empresa_id=emp.id if emp else None,
            outras_informacoes=request.form.get('outras_informacoes', ''),
        )
        try:
            db.session.add(nb)
            db.session.commit()
            flash('No-break cadastrado com sucesso!', 'success')
            return redirect(url_for('list_nobreaks'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('nobreak_form.html', proximo_codigo=gerar_codigo_interno(NoBreak, 'NBR'),
                           **opcoes_nobreak())


@app.route('/nobreaks/editar/<int:id>', methods=['GET', 'POST'])
def editar_nobreak(id):
    nb = NoBreak.query.get_or_404(id)
    if request.method == 'POST':
        nb.codigo_interno = request.form['codigo_interno']
        nb.marca = request.form['marca']
        nb.capacidade_va = request.form['capacidade_va']
        nb.potencia_watts = request.form['potencia_watts']
        emp = empresa_do_form()
        al = area_local_do_form()
        nb.local_de_uso = f'{al.area} - {al.local}' if al else ''
        nb.area_local_id = al.id if al else None
        nb.empresa_id = emp.id if emp else None
        nb.outras_informacoes = request.form.get('outras_informacoes', '')
        try:
            db.session.commit()
            flash('No-break atualizado com sucesso!', 'success')
            return redirect(url_for('list_nobreaks'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('nobreak_form.html', item=nb, **opcoes_nobreak())


@app.route('/nobreaks/excluir/<int:id>', methods=['POST'])
def excluir_nobreak(id):
    nb = NoBreak.query.get_or_404(id)
    db.session.delete(nb)
    db.session.commit()
    flash('No-break excluído com sucesso!', 'success')
    return redirect(url_for('list_nobreaks'))


# ---------- Projetores ----------

@app.route('/projetores')
def list_projetores():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(Projetor, {
        'empresa': Projetor.empresa_id,
        'local_de_uso': Projetor.local_de_uso,
        'marca': Projetor.marca,
        'tipo': Projetor.tipo,
        'resolucao': Projetor.resolucao,
    }, padrao='empresa')
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    mapa_empresas = {e.id: e.nome_fantasia for e in Empresa.query.all()}
    return render_template('projetores.html', items=items, ordem=ordem, direcao=direcao,
                           mapa_empresas=mapa_empresas, empresa_filtro=empresa_filtro)


@app.route('/projetores/novo', methods=['GET', 'POST'])
def novo_projetor():
    if request.method == 'POST':
        emp = empresa_do_form()
        al = area_local_do_form()
        prj = Projetor(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(Projetor, 'PRJ'),
            marca=request.form['marca'],
            tecnologia_imagem=request.form['tecnologia_imagem'],
            tipo=request.form['tipo'],
            resolucao=request.form['resolucao'],
            conectividade=request.form['conectividade'],
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            empresa_id=emp.id if emp else None,
            outras_informacoes=request.form.get('outras_informacoes', ''),
        )
        try:
            db.session.add(prj)
            db.session.commit()
            flash('Projetor cadastrado com sucesso!', 'success')
            return redirect(url_for('list_projetores'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('projetor_form.html', proximo_codigo=gerar_codigo_interno(Projetor, 'PRJ'))


@app.route('/projetores/editar/<int:id>', methods=['GET', 'POST'])
def editar_projetor(id):
    prj = Projetor.query.get_or_404(id)
    if request.method == 'POST':
        prj.codigo_interno = request.form['codigo_interno']
        prj.marca = request.form['marca']
        prj.tecnologia_imagem = request.form['tecnologia_imagem']
        prj.tipo = request.form['tipo']
        prj.resolucao = request.form['resolucao']
        prj.conectividade = request.form['conectividade']
        emp = empresa_do_form()
        al = area_local_do_form()
        prj.local_de_uso = f'{al.area} - {al.local}' if al else ''
        prj.area_local_id = al.id if al else None
        prj.empresa_id = emp.id if emp else None
        prj.outras_informacoes = request.form.get('outras_informacoes', '')
        try:
            db.session.commit()
            flash('Projetor atualizado com sucesso!', 'success')
            return redirect(url_for('list_projetores'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('projetor_form.html', item=prj)


@app.route('/projetores/excluir/<int:id>', methods=['POST'])
def excluir_projetor(id):
    prj = Projetor.query.get_or_404(id)
    db.session.delete(prj)
    db.session.commit()
    flash('Projetor excluído com sucesso!', 'success')
    return redirect(url_for('list_projetores'))


# ---------- Televisões ----------

@app.route('/televisoes')
def list_televisoes():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(Televisao, {
        'codigo_interno': Televisao.codigo_interno,
        'empresa': Televisao.empresa_id,
        'local_de_uso': Televisao.local_de_uso,
        'marca': Televisao.marca,
        'tipo_tecnologia_imagem': Televisao.tipo_tecnologia_imagem,
        'resolucao': Televisao.resolucao,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    mapa_empresas = {e.id: e.nome_fantasia for e in Empresa.query.all()}
    return render_template('televisoes.html', items=items, ordem=ordem, direcao=direcao,
                           mapa_empresas=mapa_empresas, empresa_filtro=empresa_filtro)


@app.route('/televisoes/novo', methods=['GET', 'POST'])
def nova_televisao():
    if request.method == 'POST':
        conectividade_video = ', '.join(request.form.getlist('conectividade_video'))
        emp = empresa_do_form()
        al = area_local_do_form()
        tv = Televisao(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(Televisao, 'TV'),
            marca=request.form['marca'],
            tipo_tecnologia_imagem=request.form['tipo_tecnologia_imagem'],
            resolucao=request.form['resolucao'],
            conectividade_video=conectividade_video,
            conectividade_rede=request.form['conectividade_rede'],
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            empresa_id=emp.id if emp else None,
            outras_informacoes=request.form.get('outras_informacoes', ''),
        )
        try:
            db.session.add(tv)
            db.session.commit()
            flash('Televisão cadastrada com sucesso!', 'success')
            return redirect(url_for('list_televisoes'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('televisao_form.html', proximo_codigo=gerar_codigo_interno(Televisao, 'TV'),
                           **opcoes_televisao())


@app.route('/televisoes/editar/<int:id>', methods=['GET', 'POST'])
def editar_televisao(id):
    tv = Televisao.query.get_or_404(id)
    if request.method == 'POST':
        tv.codigo_interno = request.form['codigo_interno']
        tv.marca = request.form['marca']
        tv.tipo_tecnologia_imagem = request.form['tipo_tecnologia_imagem']
        tv.resolucao = request.form['resolucao']
        tv.conectividade_video = ', '.join(request.form.getlist('conectividade_video'))
        tv.conectividade_rede = request.form['conectividade_rede']
        emp = empresa_do_form()
        al = area_local_do_form()
        tv.local_de_uso = f'{al.area} - {al.local}' if al else ''
        tv.area_local_id = al.id if al else None
        tv.empresa_id = emp.id if emp else None
        tv.outras_informacoes = request.form.get('outras_informacoes', '')
        try:
            db.session.commit()
            flash('Televisão atualizada com sucesso!', 'success')
            return redirect(url_for('list_televisoes'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('televisao_form.html', item=tv, **opcoes_televisao())


@app.route('/televisoes/excluir/<int:id>', methods=['POST'])
def excluir_televisao(id):
    tv = Televisao.query.get_or_404(id)
    db.session.delete(tv)
    db.session.commit()
    flash('Televisão excluída com sucesso!', 'success')
    return redirect(url_for('list_televisoes'))


# ---------- Racks ----------

@app.route('/racks')
def list_racks():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(Rack, {
        'codigo_interno': Rack.codigo_interno,
        'empresa': Rack.empresa_id,
        'local_de_uso': Rack.local_de_uso,
        'marca': Rack.marca,
        'modelo': Rack.modelo,
        'us': Rack.us,
    })
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    mapa_empresas = {e.id: e.nome_fantasia for e in Empresa.query.all()}
    return render_template('racks.html', items=items, ordem=ordem, direcao=direcao,
                           mapa_empresas=mapa_empresas, empresa_filtro=empresa_filtro)


@app.route('/racks/novo', methods=['GET', 'POST'])
def novo_rack():
    if request.method == 'POST':
        emp = empresa_do_form()
        al = area_local_do_form()
        r = Rack(
            codigo_interno=request.form.get('codigo_interno', '').strip() or gerar_codigo_interno(Rack, 'RCK'),
            marca=request.form['marca'],
            cor=request.form['cor'],
            modelo=request.form['modelo'],
            largura=request.form['largura'],
            profundidade=request.form['profundidade'],
            altura=request.form['altura'],
            us=request.form['us'],
            peso_produto=request.form['peso_produto'],
            capacidade_carga=request.form.get('capacidade_carga', ''),
            fechadura=request.form['fechadura'],
            classificacao_ip=request.form.get('classificacao_ip', ''),
            pes_niveladores=request.form['pes_niveladores'],
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            empresa_id=emp.id if emp else None,
            outras_informacoes=request.form.get('outras_informacoes', ''),
        )
        try:
            db.session.add(r)
            db.session.commit()
            flash('Rack cadastrado com sucesso!', 'success')
            return redirect(url_for('list_racks'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('rack_form.html', proximo_codigo=gerar_codigo_interno(Rack, 'RCK'),
                           **opcoes_rack())


@app.route('/racks/editar/<int:id>', methods=['GET', 'POST'])
def editar_rack(id):
    r = Rack.query.get_or_404(id)
    if request.method == 'POST':
        r.codigo_interno = request.form['codigo_interno']
        r.marca = request.form['marca']
        r.cor = request.form['cor']
        r.modelo = request.form['modelo']
        r.largura = request.form['largura']
        r.profundidade = request.form['profundidade']
        r.altura = request.form['altura']
        r.us = request.form['us']
        r.peso_produto = request.form['peso_produto']
        r.capacidade_carga = request.form.get('capacidade_carga', '')
        r.fechadura = request.form['fechadura']
        r.classificacao_ip = request.form.get('classificacao_ip', '')
        r.pes_niveladores = request.form['pes_niveladores']
        emp = empresa_do_form()
        al = area_local_do_form()
        r.local_de_uso = f'{al.area} - {al.local}' if al else ''
        r.area_local_id = al.id if al else None
        r.empresa_id = emp.id if emp else None
        r.outras_informacoes = request.form.get('outras_informacoes', '')
        try:
            db.session.commit()
            flash('Rack atualizado com sucesso!', 'success')
            return redirect(url_for('list_racks'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('rack_form.html', item=r, **opcoes_rack())


@app.route('/racks/excluir/<int:id>', methods=['POST'])
def excluir_rack(id):
    r = Rack.query.get_or_404(id)
    db.session.delete(r)
    db.session.commit()
    flash('Rack excluído com sucesso!', 'success')
    return redirect(url_for('list_racks'))


# ---------- Softwares ----------

@app.route('/softwares')
def list_softwares():
    categoria = request.args.get('categoria', 'all')
    empresa_filtro = request.args.get('empresa', 'all')
    colunas = {
        'codigo_equipamento': Software.codigo_equipamento,
        'categoria': Software.categoria,
        'nome': Software.nome,
        'marca': Software.marca,
        'tipo_licenciamento': Software.tipo_licenciamento,
        'validade_assinatura': Software.validade_assinatura,
    }
    ordem = request.args.get('ordem', 'nome')
    direcao = request.args.get('direcao', 'asc')
    if ordem not in colunas:
        ordem = 'nome'
    if direcao not in ('asc', 'desc'):
        direcao = 'asc'
    coluna = colunas[ordem].desc() if direcao == 'desc' else colunas[ordem]
    consulta = Software.query
    if categoria != 'all':
        consulta = consulta.filter_by(categoria=categoria)
    items = consulta.order_by(coluna).all()
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('softwares.html', items=items, categoria=categoria, ordem=ordem,
                           direcao=direcao, empresa_filtro=empresa_filtro)


@app.route('/softwares/novo', methods=['GET', 'POST'])
def novo_software():
    categorias = [
        'Sistema Operacional', 'Banco de dados', 'Sistema de Gestão',
        'Antivírus', 'Software de Escritório', 'Software de Desenvolvimento',
        'Software de Geoprocessamento', 'Software de Inteligência Artificial',
        'Software Utilitário', 'Software Geral'
    ]
    if request.method == 'POST':
        validade = request.form.get('validade_assinatura', '')
        emp = empresa_do_form()
        al = area_local_do_form()
        valores = [v for v in request.form.getlist('codigo_equipamento') if v]
        codigo_eq, w_ids, n_ids, a_ids = processar_equipamentos_software(valores)
        sw = Software(
            codigo_equipamento=codigo_eq,
            categoria=request.form['categoria'],
            marca=request.form['marca'],
            nome=request.form['nome'],
            descricao=request.form.get('descricao', ''),
            tipo_licenciamento=request.form['tipo_licenciamento'].strip(),
            validade_assinatura=date.fromisoformat(validade) if validade else None,
            chave=request.form.get('chave', ''),
            local_de_uso=(f'{al.area} - {al.local}' if al else ''),
            area_local_id=al.id if al else None,
            empresa_id=emp.id if emp else None,
            workstation_id=w_ids[0] if w_ids else None,
            notebook_id=n_ids[0] if n_ids else None,
            outras_informacoes=request.form.get('outras_informacoes', ''),
        )
        try:
            db.session.add(sw)
            db.session.flush()
            sincronizar_vinculos_software(sw, w_ids, n_ids, a_ids)
            db.session.commit()
            flash('Software cadastrado com sucesso!', 'success')
            return redirect(url_for('list_softwares'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('software_form.html', categorias=categorias, equipamentos_selecionados='',
                           tipo_licenciamento='', **opcoes_licencas_software())


@app.route('/softwares/editar/<int:id>', methods=['GET', 'POST'])
def editar_software(id):
    categorias = [
        'Sistema Operacional', 'Banco de dados', 'Sistema de Gestão',
        'Antivírus', 'Software de Escritório', 'Software de Desenvolvimento',
        'Software de Geoprocessamento', 'Software de Inteligência Artificial',
        'Software Utilitário', 'Software Geral'
    ]
    sw = Software.query.get_or_404(id)
    if request.method == 'POST':
        validade = request.form.get('validade_assinatura', '')
        valores = [v for v in request.form.getlist('codigo_equipamento') if v]
        codigo_eq, w_ids, n_ids, a_ids = processar_equipamentos_software(valores)
        sw.codigo_equipamento = codigo_eq
        sw.workstation_id = w_ids[0] if w_ids else None
        sw.notebook_id = n_ids[0] if n_ids else None
        sw.categoria = request.form['categoria']
        sw.marca = request.form['marca']
        sw.nome = request.form['nome']
        sw.descricao = request.form.get('descricao', '')
        sw.tipo_licenciamento = request.form['tipo_licenciamento'].strip()
        sw.validade_assinatura = date.fromisoformat(validade) if validade else None
        sw.chave = request.form.get('chave', '')
        emp = empresa_do_form()
        al = area_local_do_form()
        sw.local_de_uso = f'{al.area} - {al.local}' if al else ''
        sw.area_local_id = al.id if al else None
        sw.empresa_id = emp.id if emp else None
        sw.outras_informacoes = request.form.get('outras_informacoes', '')
        try:
            sincronizar_vinculos_software(sw, w_ids, n_ids, a_ids)
            db.session.commit()
            flash('Software atualizado com sucesso!', 'success')
            return redirect(url_for('list_softwares'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('software_form.html', item=sw, categorias=categorias,
                           tipo_licenciamento=sw.tipo_licenciamento,
                           equipamentos_selecionados=equipamentos_selecionados(sw),
                           **opcoes_licencas_software())


@app.route('/softwares/excluir/<int:id>', methods=['POST'])
def excluir_software(id):
    sw = Software.query.get_or_404(id)
    SoftwareEquipamento.query.filter_by(software_id=sw.id).delete()
    db.session.delete(sw)
    db.session.commit()
    flash('Software excluído com sucesso!', 'success')
    return redirect(url_for('list_softwares'))


# ---------- Empresas ----------

@app.route('/empresas')
def list_empresas():
    items, ordem, direcao = ordenar_lista(Empresa, {
        'cnpj': Empresa.cnpj,
        'nome_fantasia': Empresa.nome_fantasia,
        'telefones': Empresa.telefones,
        'email': Empresa.email,
    }, padrao='cnpj')
    return render_template('empresas.html', items=items, ordem=ordem, direcao=direcao)


@app.route('/empresas/novo', methods=['GET', 'POST'])
def nova_empresa():
    if request.method == 'POST':
        emp = Empresa(
            razao_social=request.form['razao_social'],
            nome_fantasia=request.form['nome_fantasia'],
            cnpj=request.form['cnpj'],
            cep=request.form['cep'],
            logradouro=request.form['logradouro'],
            bairro=request.form['bairro'],
            cidade=request.form['cidade'],
            estado=request.form['estado'],
            telefones=request.form['telefones'],
            email=request.form['email'],
            outras_informacoes=request.form.get('outras_informacoes', ''),
        )
        try:
            db.session.add(emp)
            db.session.commit()
            flash('Empresa cadastrada com sucesso!', 'success')
            return redirect(url_for('list_empresas'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('empresa_form.html')


@app.route('/empresas/editar/<int:id>', methods=['GET', 'POST'])
def editar_empresa(id):
    emp = Empresa.query.get_or_404(id)
    if request.method == 'POST':
        emp.razao_social = request.form['razao_social']
        emp.nome_fantasia = request.form['nome_fantasia']
        emp.cnpj = request.form['cnpj']
        emp.cep = request.form['cep']
        emp.logradouro = request.form['logradouro']
        emp.bairro = request.form['bairro']
        emp.cidade = request.form['cidade']
        emp.estado = request.form['estado']
        emp.telefones = request.form['telefones']
        emp.email = request.form['email']
        emp.outras_informacoes = request.form.get('outras_informacoes', '')
        try:
            db.session.commit()
            flash('Empresa atualizada com sucesso!', 'success')
            return redirect(url_for('list_empresas'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('empresa_form.html', item=emp)


@app.route('/empresas/excluir/<int:id>', methods=['POST'])
def excluir_empresa(id):
    emp = Empresa.query.get_or_404(id)
    for modelo in (Workstation, EstacaoAllInOne, Notebook, Monitor, Periferico, Impressora, Tablet, Servidor, Projetor, Televisao, Modem, Roteador,
                   Switch, Firewall, AccessPoint, NoBreak, Rack, Software, Usuario, AreaLocal):
        modelo.query.filter_by(empresa_id=emp.id).update({'empresa_id': None})
    db.session.delete(emp)
    db.session.commit()
    flash('Empresa excluída com sucesso!', 'success')
    return redirect(url_for('list_empresas'))


# ---------- Usuários ----------

@app.route('/usuarios')
def list_usuarios():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(Usuario, {
        'cpf': Usuario.cpf,
        'nome_completo': Usuario.nome_completo,
        'empresa_que_trabalha': Usuario.empresa_que_trabalha,
        'telefones': Usuario.telefones,
        'email': Usuario.email,
    }, padrao='nome_completo')
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('usuarios.html', items=items, ordem=ordem, direcao=direcao,
                           empresa_filtro=empresa_filtro)


@app.route('/usuarios/novo', methods=['GET', 'POST'])
def novo_usuario():
    if request.method == 'POST':
        emp = empresa_do_form()
        usr = Usuario(
            cpf=request.form['cpf'],
            empresa_que_trabalha=emp.nome_fantasia if emp else '',
            empresa_id=emp.id if emp else None,
            nome_completo=request.form['nome_completo'],
            cep=request.form['cep'],
            logradouro=request.form['logradouro'],
            bairro=request.form['bairro'],
            cidade=request.form['cidade'],
            estado=request.form['estado'],
            telefones=request.form['telefones'],
            email=request.form['email'],
            outras_informacoes=request.form.get('outras_informacoes', ''),
        )
        try:
            db.session.add(usr)
            db.session.commit()
            flash('Usuário cadastrado com sucesso!', 'success')
            return redirect(url_for('list_usuarios'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('usuario_form.html')


@app.route('/usuarios/editar/<int:id>', methods=['GET', 'POST'])
def editar_usuario(id):
    usr = Usuario.query.get_or_404(id)
    if request.method == 'POST':
        emp = empresa_do_form()
        usr.cpf = request.form['cpf']
        usr.empresa_que_trabalha = emp.nome_fantasia if emp else ''
        usr.empresa_id = emp.id if emp else None
        usr.nome_completo = request.form['nome_completo']
        usr.cep = request.form['cep']
        usr.logradouro = request.form['logradouro']
        usr.bairro = request.form['bairro']
        usr.cidade = request.form['cidade']
        usr.estado = request.form['estado']
        usr.telefones = request.form['telefones']
        usr.email = request.form['email']
        usr.outras_informacoes = request.form.get('outras_informacoes', '')
        try:
            db.session.commit()
            flash('Usuário atualizado com sucesso!', 'success')
            return redirect(url_for('list_usuarios'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('usuario_form.html', item=usr)


@app.route('/usuarios/excluir/<int:id>', methods=['POST'])
def excluir_usuario(id):
    usr = Usuario.query.get_or_404(id)
    Workstation.query.filter_by(usuario_id=usr.id).update({'usuario_id': None})
    EstacaoAllInOne.query.filter_by(usuario_id=usr.id).update({'usuario_id': None})
    Notebook.query.filter_by(usuario_id=usr.id).update({'usuario_id': None})
    Monitor.query.filter_by(usuario_id=usr.id).update({'usuario_id': None})
    Periferico.query.filter_by(usuario_id=usr.id).update({'usuario_id': None})
    Tablet.query.filter_by(usuario_id=usr.id).update({'usuario_id': None})
    db.session.delete(usr)
    db.session.commit()
    flash('Usuário excluído com sucesso!', 'success')
    return redirect(url_for('list_usuarios'))


# ---------- Administradores ----------

@app.route('/administradores')
def list_administradores():
    items, ordem, direcao = ordenar_lista(Administrador, {
        'cpf': Administrador.cpf,
        'nome_completo': Administrador.nome_completo,
        'telefone': Administrador.telefone,
        'email': Administrador.email,
        'usuario': Administrador.usuario,
    }, padrao='cpf')
    return render_template('administradores.html', items=items, ordem=ordem, direcao=direcao)


@app.route('/administradores/novo', methods=['GET', 'POST'])
def novo_administrador():
    if request.method == 'POST':
        senha_inicial = senha_inicial_administrador_atual()
        if not senha_inicial:
            flash('Configure a Senha Inicial de um Usuário Administrador antes de cadastrar administradores.', 'error')
            return redirect(url_for('configuracoes_sistema'))
        adm = Administrador(
            cpf=request.form['cpf'],
            nome_completo=request.form['nome_completo'],
            telefone=request.form['telefone'],
            email=request.form['email'],
            usuario=request.form['usuario'],
            tipo=request.form['tipo'],
            recebe_notificacoes=request.form['recebe_notificacoes'],
            senha=md5(senha_inicial),
            outras_informacoes=request.form.get('outras_informacoes', ''),
        )
        try:
            db.session.add(adm)
            db.session.commit()
            flash('Administrador cadastrado com sucesso!', 'success')
            return redirect(url_for('list_administradores'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('administrador_form.html')


@app.route('/administradores/editar/<int:id>', methods=['GET', 'POST'])
def editar_administrador(id):
    adm = Administrador.query.get_or_404(id)
    if request.method == 'POST':
        adm.cpf = request.form['cpf']
        adm.nome_completo = request.form['nome_completo']
        adm.telefone = request.form['telefone']
        adm.email = request.form['email']
        adm.usuario = request.form['usuario']
        adm.tipo = request.form['tipo']
        adm.recebe_notificacoes = request.form['recebe_notificacoes']
        adm.outras_informacoes = request.form.get('outras_informacoes', '')
        try:
            db.session.commit()
            flash('Administrador atualizado com sucesso!', 'success')
            return redirect(url_for('list_administradores'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('administrador_form.html', item=adm)


@app.route('/administradores/excluir/<int:id>', methods=['POST'])
def excluir_administrador(id):
    adm = Administrador.query.get_or_404(id)
    db.session.delete(adm)
    db.session.commit()
    flash('Administrador excluído com sucesso!', 'success')
    return redirect(url_for('list_administradores'))


@app.route('/administradores/resetar-senha/<int:id>', methods=['POST'])
def resetar_senha_administrador(id):
    adm = Administrador.query.get_or_404(id)
    smtp = ConfiguracaoSistema.query.order_by(ConfiguracaoSistema.id.asc()).first()
    if smtp and smtp.usuario_autenticacao and smtp.senha_autenticacao:
        nova_senha = senha_aleatoria()
        adm.senha = md5(nova_senha)
        db.session.commit()
        try:
            enviar_email_senha(adm, nova_senha, tipo='redefinicao')
            flash('Senha redefinida e e-mail enviado com sucesso!', 'success')
        except Exception as erro:
            flash(f'Senha redefinida, mas falha ao enviar o e-mail: {erro}', 'error')
    else:
        flash('Serviço de e-mail não configurado; a senha não foi alterada.', 'error')
    return redirect(url_for('list_administradores'))


# ---------- Configurações do sistema ----------

PORTAS_SMTP = (25, 465, 587, 995, 2525)


def configuracao_sistema_para_formulario():
    configuracao = ConfiguracaoSistema.query.order_by(ConfiguracaoSistema.id.asc()).first()
    if configuracao:
        return configuracao
    return ConfiguracaoSistema(
        nome_fantasia_empresa='',
        url_acesso_sistema='',
        senha_inicial_administrador='',
        servidor_smtp='',
        porta_smtp=None,
        usuario_autenticacao='',
        senha_autenticacao='',
        nome_remetente='',
        assunto_esqueci_senha='',
        mensagem_esqueci_senha='',
        assunto_redefinicao_senha='',
        mensagem_redefinicao_senha='',
        assunto_software_assinatura=ASSUNTO_SOFTWARE_ASSINATURA_PADRAO,
        mensagem_software_assinatura=MENSAGEM_SOFTWARE_ASSINATURA_PADRAO,
    )


def senha_inicial_administrador_atual():
    configuracao = ConfiguracaoSistema.query.order_by(ConfiguracaoSistema.id.asc()).first()
    if configuracao and configuracao.senha_inicial_administrador.strip():
        return configuracao.senha_inicial_administrador
    return None


@app.route('/configuracoes-sistema', methods=['GET', 'POST'])
def configuracoes_sistema():
    configuracao = configuracao_sistema_para_formulario()
    if request.method == 'POST':
        nome_fantasia = request.form.get('nome_fantasia_empresa', '').strip()
        url_acesso = request.form.get('url_acesso_sistema', '').strip()
        senha_inicial_administrador = request.form.get('senha_inicial_administrador', '')
        servidor_smtp = request.form.get('servidor_smtp', '').strip()
        porta_recebida = request.form.get('porta_smtp', '')
        usuario = request.form.get('usuario_autenticacao', '').strip()
        senha = request.form.get('senha_autenticacao', '')
        remetente = request.form.get('nome_remetente', '').strip()
        assunto_esqueci = request.form.get('assunto_esqueci_senha', '').strip()
        mensagem_esqueci = request.form.get('mensagem_esqueci_senha', '')
        assunto_redefinicao = request.form.get('assunto_redefinicao_senha', '').strip()
        mensagem_redefinicao = request.form.get('mensagem_redefinicao_senha', '')
        assunto_software_assinatura = request.form.get('assunto_software_assinatura', '').strip()
        mensagem_software_assinatura = request.form.get('mensagem_software_assinatura', '')
        porta = int(porta_recebida) if porta_recebida in {str(p) for p in PORTAS_SMTP} else None

        configuracao.nome_fantasia_empresa = nome_fantasia
        configuracao.url_acesso_sistema = url_acesso
        configuracao.senha_inicial_administrador = senha_inicial_administrador
        configuracao.servidor_smtp = servidor_smtp
        configuracao.porta_smtp = porta
        configuracao.usuario_autenticacao = usuario
        configuracao.senha_autenticacao = senha
        configuracao.nome_remetente = remetente
        configuracao.assunto_esqueci_senha = assunto_esqueci
        configuracao.mensagem_esqueci_senha = mensagem_esqueci
        configuracao.assunto_redefinicao_senha = assunto_redefinicao
        configuracao.mensagem_redefinicao_senha = mensagem_redefinicao
        configuracao.assunto_software_assinatura = assunto_software_assinatura
        configuracao.mensagem_software_assinatura = mensagem_software_assinatura

        campos_obrigatorios = (nome_fantasia, url_acesso, senha_inicial_administrador.strip(), servidor_smtp, porta,
                               usuario, senha, remetente, assunto_esqueci, mensagem_esqueci.strip(),
                               assunto_redefinicao, mensagem_redefinicao.strip(), assunto_software_assinatura,
                               mensagem_software_assinatura.strip())
        limites = ((nome_fantasia, 100), (url_acesso, 200), (senha_inicial_administrador, 50),
                   (servidor_smtp, 200),
                   (usuario, 200), (senha, 200), (remetente, 100),
                   (assunto_esqueci, 100), (assunto_redefinicao, 100), (assunto_software_assinatura, 100))
        if not all(campos_obrigatorios):
            flash('Preencha todos os campos obrigatórios.', 'error')
        elif any(len(valor) > limite for valor, limite in limites):
            flash('Um ou mais campos excedem o limite de caracteres permitido.', 'error')
        else:
            try:
                if configuracao.id is None:
                    db.session.add(configuracao)
                db.session.commit()
                flash('Configurações salvas com sucesso.', 'success')
                return redirect(url_for('configuracoes_sistema'))
            except Exception as erro:
                db.session.rollback()
                flash(f'Erro ao salvar as configurações: {erro}', 'error')
                configuracao = configuracao_sistema_para_formulario()

    return render_template(
        'configuracoes_sistema_form.html',
        configuracao=configuracao,
        portas_smtp=PORTAS_SMTP,
    )


# ---------- Chaves de Acesso ----------

@app.route('/chaves-acesso')
def list_chaves_acesso():
    items, ordem, direcao = ordenar_lista(ChaveAcesso, {
        'cpf': ChaveAcesso.cpf,
        'nome_completo': ChaveAcesso.nome_completo,
        'telefone': ChaveAcesso.telefone,
        'email': ChaveAcesso.email,
    }, padrao='cpf')
    return render_template('chaves_acesso.html', items=items, ordem=ordem, direcao=direcao)


@app.route('/chaves-acesso/novo', methods=['GET', 'POST'])
def nova_chave_acesso():
    if request.method == 'POST':
        necessidade = request.form.get('necessidade', '')
        if necessidade not in NECESSIDADES_CHAVE_ACESSO:
            flash('Selecione uma opção válida para Necessidade.', 'error')
            return render_template('chave_acesso_form.html', necessidades=NECESSIDADES_CHAVE_ACESSO)
        item = ChaveAcesso(
            cpf=request.form['cpf'],
            nome_completo=request.form['nome_completo'],
            telefone=request.form['telefone'],
            email=request.form['email'],
            cep=request.form['cep'],
            logradouro=request.form['logradouro'],
            bairro=request.form['bairro'],
            cidade=request.form['cidade'],
            estado=request.form['estado'],
            necessidade=necessidade,
            necessidade_webservices=request.form['necessidade_webservices'],
            chave_acesso=gerar_chave_acesso(),
        )
        try:
            db.session.add(item)
            db.session.commit()
            flash('Chave de Acesso cadastrada com sucesso!', 'success')
            return redirect(url_for('list_chaves_acesso'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('chave_acesso_form.html', necessidades=NECESSIDADES_CHAVE_ACESSO)


@app.route('/chaves-acesso/editar/<int:id>', methods=['GET', 'POST'])
def editar_chave_acesso(id):
    item = ChaveAcesso.query.get_or_404(id)
    if request.method == 'POST':
        necessidade = request.form.get('necessidade', '')
        if necessidade not in NECESSIDADES_CHAVE_ACESSO:
            flash('Selecione uma opção válida para Necessidade.', 'error')
            return render_template(
                'chave_acesso_form.html', item=item,
                necessidades=NECESSIDADES_CHAVE_ACESSO
            )
        item.cpf = request.form['cpf']
        item.nome_completo = request.form['nome_completo']
        item.telefone = request.form['telefone']
        item.email = request.form['email']
        item.cep = request.form['cep']
        item.logradouro = request.form['logradouro']
        item.bairro = request.form['bairro']
        item.cidade = request.form['cidade']
        item.estado = request.form['estado']
        item.necessidade = necessidade
        item.necessidade_webservices = request.form['necessidade_webservices']
        try:
            db.session.commit()
            flash('Chave de Acesso atualizada com sucesso!', 'success')
            return redirect(url_for('list_chaves_acesso'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template(
        'chave_acesso_form.html', item=item,
        necessidades=NECESSIDADES_CHAVE_ACESSO
    )


@app.route('/chaves-acesso/excluir/<int:id>', methods=['POST'])
def excluir_chave_acesso(id):
    item = ChaveAcesso.query.get_or_404(id)
    try:
        db.session.delete(item)
        db.session.commit()
        flash('Chave de Acesso excluída com sucesso!', 'success')
    except Exception as e:
        db.session.rollback()
        flash(f'Erro ao excluir: {str(e)}', 'error')
    return redirect(url_for('list_chaves_acesso'))


# ---------- Webservices ----------

@app.route('/webservices')
def webservices():
    return render_template('webservices.html')


# ---------- Áreas e Locais ----------

@app.route('/areas-locais')
def list_areas_locais():
    empresa_filtro = request.args.get('empresa', 'all')
    items, ordem, direcao = ordenar_lista(AreaLocal, {
        'empresa': AreaLocal.empresa,
        'area': AreaLocal.area,
        'local': AreaLocal.local,
    }, padrao='empresa')
    if empresa_filtro != 'all':
        items = [item for item in items if str(item.empresa_id) == empresa_filtro]
    return render_template('areas_locais.html', items=items, ordem=ordem, direcao=direcao,
                           empresa_filtro=empresa_filtro)


@app.route('/areas-locais/novo', methods=['GET', 'POST'])
def nova_area_local():
    if request.method == 'POST':
        emp = empresa_do_form()
        al = AreaLocal(
            empresa=emp.nome_fantasia if emp else '',
            empresa_id=emp.id if emp else None,
            area=request.form['area'],
            local=request.form['local'],
            outras_informacoes=request.form.get('outras_informacoes', ''),
        )
        try:
            db.session.add(al)
            db.session.commit()
            flash('Área/Local cadastrado com sucesso!', 'success')
            return redirect(url_for('list_areas_locais'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao cadastrar: {str(e)}', 'error')
    return render_template('area_local_form.html')


@app.route('/areas-locais/editar/<int:id>', methods=['GET', 'POST'])
def editar_area_local(id):
    al = AreaLocal.query.get_or_404(id)
    if request.method == 'POST':
        emp = empresa_do_form()
        al.empresa = emp.nome_fantasia if emp else ''
        al.empresa_id = emp.id if emp else None
        al.area = request.form['area']
        al.local = request.form['local']
        al.outras_informacoes = request.form.get('outras_informacoes', '')
        try:
            db.session.commit()
            flash('Área/Local atualizado com sucesso!', 'success')
            return redirect(url_for('list_areas_locais'))
        except Exception as e:
            db.session.rollback()
            flash(f'Erro ao atualizar: {str(e)}', 'error')
    return render_template('area_local_form.html', item=al)


@app.route('/areas-locais/excluir/<int:id>', methods=['POST'])
def excluir_area_local(id):
    al = AreaLocal.query.get_or_404(id)
    for modelo in (Workstation, EstacaoAllInOne, Notebook, Monitor, Periferico, Impressora, Tablet, Servidor, Projetor, Televisao, Modem, Roteador,
                   Switch, Firewall, AccessPoint, NoBreak, Rack, Software):
        modelo.query.filter_by(area_local_id=al.id).update({'area_local_id': None})
    db.session.delete(al)
    db.session.commit()
    flash('Área/Local excluído com sucesso!', 'success')
    return redirect(url_for('list_areas_locais'))


if __name__ == '__main__':
    with app.app_context():
        db.create_all()
        migrar_configuracoes_email_software()
        converter_senhas_para_md5()
        migrar_vinculos_softwares()
    app.run(debug=True, port=5000)
